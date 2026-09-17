#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""stepgate — 보고자/심판자 분리 단계 게이트 (hook 차단 환경의 forcing function).

세 역할을 분리한다:
  - 워커(스킬 실행 모델): `form`으로 *관찰 항목만* 받아 정직히 답한다. **합격 기준은 모른다.**
      · 표본 추출 없음 — **모든 관찰 항목이 매 회차 전부 출제**된다(기준이 그 회차에 빠지지 않음).
      · ★ 극성 회전(`variants`): 같은 항목을 회차마다 *Y정답 표현* 또는 *N정답 표현* 중 하나로 물어
        **고정 정답 벡터 암기를 무력화**한다(예: "①②③을 다 적었는가"=Y ↔ "빠진 것이 있는가"=N).
        뽑힌 변형 번호는 form 상태에 기록되고 `judge`가 *그 변형의 must*로 채점 → `form` 선행 필수.
        FAIL이면 출제를 폐기해 재시도 시 극성이 새로 굴러간다(답만 뒤집는 시행착오 차단).
  - 결정론 심판(이 스크립트의 `judge`): 워커가 안 보는 **봉인 루브릭**을 읽어 PASS/FAIL.
      · 디스크로 확인 가능한 항목은 워커 도장을 무시하고 **직접 재검증**(거짓 차단).
      · attest(인지) 항목은 워커 답을 받되 '미검증' 표기 — 단 rubric의 attest_must와 다르면 FAIL
        (예: '직전 폴더 열었나?'에 정직히 N이면 스스로 스킵을 시인 → 차단).
      · 직전 체크포인트 통과 토큰을 확인(순서 강제).
  - substance 심판(레거시 rubric escalate=true): judge가 NEEDS-JUDGE를 내면 워커가 *별도 judge
      서브에이전트*를 띄워 산출물 실물을 읽혀 판정시키고, subverdict=PASS|FAIL로 확정.
  - 비동기 빡빡 감사(rubric "mode":"async" — 신규 기본): judge가 바닥(disk+attest+quote)만 즉시
      판정하고 **escalate를 인라인으로 막지 않는다**. 대신 감사 큐에 적재(enqueue)하고 워커는 다음
      작업으로 넘어간다(중첩). 별도 백그라운드 gate-audit Agent가 audit_rubrics/<skill>.json의 전
      lens로 산출물을 적대적 검수(`audit-judge`)하고, 발행 직전 `resolve`가 전 감사 PASS를 확인(T3).

동시 워크플로 격리(★ 2026-06-29): 모든 영속 상태(ledger 순서토큰·감사 큐·form 출제)를 *skill* 단위가
  아니라 **skill + 워크플로 키(wf)** 단위로 파티션한다. wf는 인자에서 자동 도출 —
  명시 wf= > dir=(과제폴더) 해시 > deliverable 부모폴더 해시 > .wfid 폴백(단일작업 호환).
  같은 과제폴더의 동시 작업은 같은 wf(한 워크플로)로 보되, *산출물 구분*은 deliv_hash로 한다
  (같은 cp라도 다른 산출물끼리 SUPERSEDE/HOLD로 서로 죽이지 않음). 다른 폴더 = 다른 wf = 완전 격리.

봉인 루브릭: 엔진과 같은 폴더의 rubrics/<skill>.json (_stepgate/rubrics/). **대상 스킬 SKILL.md는 이 경로를
  참조하지 않는다**(봉인 = 워커 컨텍스트에 안 들어옴). 임계값·합격조건은 여기에만 둔다.
  강제력은 *비밀에 의존하지 않게* — 검증 가능 항목은 항상 디스크 재검증, 비밀은 그 위에 가산.

verdict 토큰: %TEMP%\\claude_stepgate\\<skill>.ledger 에 최신 1개(순서 강제용). 종료 시 `clear`.

사용: (동시 워크플로 격리용 파티션 인자 — 모든 단계에 일관되게: wf=<키> 또는 dir=<과제폴더>;
       judge/enqueue/resolve는 deliverable=<abspath>만 줘도 그 부모폴더로 자동 파티션)
  stepgate.py init   <skill> [dir=<과제폴더>|wf=<키>]
  stepgate.py form   <skill> <cp> [dir=<과제폴더>|deliverable=<abspath>|wf=<키>]
  stepgate.py judge  <skill> <cp> [<id>=<값> ...] [async=1] [deliverable=<abspath>] [dir=<과제폴더>] [wf=<키>] [subverdict=PASS|FAIL] [evidence=<인용요지>]
  stepgate.py status <skill> [dir=…|deliverable=…|wf=…]
  stepgate.py clear  <skill> [dir=…|wf=…]
  stepgate.py audit  <session.jsonl> | --latest [basedir]   # 사후 룰베이스 감사
  # ── 비동기 감사(T2/T3) ──
  stepgate.py enqueue      <skill> <cp> deliverable=<abspath> [dir=…]   # 감사 큐 적재 + 발사 프롬프트
  stepgate.py audit-judge  <reqid>                                      # 감사 패킷(lens) 출력
  stepgate.py audit-judge  <reqid> --record subverdict=PASS|FAIL findings=… evidence=…
  stepgate.py resolve      <skill> deliverable=<abspath>               # T3 사인오프 게이트(전 감사 PASS 확인)
  stepgate.py audit-status <skill>                                      # 활성 감사 PENDING/PASS/FAIL 집계

exit: PASS=0 · FAIL=1 · NEEDS-JUDGE/HOLD(보류)=2 · 사용오류=3.
audit: 끝난 세션 JSONL을 *결정론 규칙*으로 스캔(모델 자의 해석 0) — 게이트 미호출(R1)·봉인 정답지
  열람(R2)·FAIL 무시(R3)·subverdict 위조(R4)·순서 위반(R5)·비동기감사 미완료(R6)·해소게이트
  누락(R7)을 탐지. AUDIT PASS=0 / FAIL=1. (audit_rubrics는 비봉인이라 R2 제외.)
주의: 인자에 cmd 메타문자(| & > <)를 넣지 말 것(python.cmd 셔임에서 깨짐).
"""
import sys, os, json, time, glob, re, random, hashlib, hmac, binascii

# Windows 한국어 콘솔(cp949) 특수문자 출력 크래시 방지 — utf-8로 강제.
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

RUBRIC_DIR = os.path.normpath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "rubrics"))


def gate_dir():
    base = os.environ.get("TEMP") or os.environ.get("TMP") or os.path.expanduser("~")
    d = os.path.join(base, "claude_stepgate")
    os.makedirs(d, exist_ok=True)
    return d


def ledger_path(skill, wf=""):
    key = skill + ("__" + wf if wf else "")
    safe = "".join(c if (c.isalnum() or c in "-_.") else "_" for c in key)
    return os.path.join(gate_dir(), safe + ".ledger")


def parse_kv(args):
    spec = {}
    for a in args:
        if "=" in a:
            k, v = a.split("=", 1)
            spec[k.strip()] = v
    return spec


def load_rubric(skill):
    p = os.path.join(RUBRIC_DIR, skill + ".json")
    if not os.path.isfile(p):
        return None
    try:
        with open(p, encoding="utf-8") as fh:
            return json.load(fh)
    except (ValueError, OSError) as e:
        print("GATE ERROR: 루브릭 로드 실패 %s :: %s" % (p, e))
        sys.exit(3)


def resolve(val, args):
    """루브릭의 {id} 자리표시자를 워커가 제출한 args[id](경로/값)로 치환.
    예: "{dir}\\_raw_sources\\*" + args[dir] → 실제 glob. 미제출 id는 리터럴로 남아 검사 실패."""
    if not isinstance(val, str):
        return val
    out = val
    for k, v in args.items():
        out = out.replace("{%s}" % k, str(v))
    return out


def resolve_re(val, args, raw=False):
    """`pat`용 치환 — 기본적으로 **치환값을 정규식으로 읽지 않는다**(`re.escape`).

    워커가 낸 값을 그대로 정규식에 꽂으면 양쪽으로 깨진다: ① `.*`처럼 메타문자가 든 값을
    내면 **아무 내용이나 통과**하고, ② `a.b`·`(주)`처럼 점·괄호가 든 *정직한* 값은
    엉뚱한 곳에 히트해 **오탐**이 난다. 치환값은 '찾을 문자열'이지 패턴이 아니다.
    치환값 자체가 패턴이어야 하는 루브릭만 `"raw_pat": true`로 명시한다.
    """
    if not isinstance(val, str):
        return val
    out = val
    for k, v in args.items():
        out = out.replace("{%s}" % k, str(v) if raw else re.escape(str(v)))
    return out


def _int_field(c, args, key, default):
    """정수 필드도 자리표시자를 받는다 — 그리고 **못 읽으면 명시적 실패**다. (값, 오류).

    예전에는 `int(c.get("min", 1))`이라 `"min": "{roster}"` 같은 기준이 미포착
    ValueError로 프로세스를 죽였다. GATE FAIL 줄조차 안 남으므로 사후 감사(R3)도
    'FAIL 무시'를 탐지할 수 없었다 — **죽은 기준이 조용히 통과처럼 보인다.**
    """
    raw = c.get(key, default)
    try:
        return int(resolve(raw, args)), None
    except (TypeError, ValueError):
        return None, "정수 필드 '%s'를 읽을 수 없다: %r (자리표시자 미제출?)" % (key, raw)


def run_criterion(c, args):
    """디스크 진실을 *직접* 재도출 (워커 도장 무시). (ok, msg).

    ★ 어떤 예외도 **명시적 FAIL로 바꾼다.** 미포착 예외가 프로세스를 죽이면
    '기준이 막았다'와 '기준이 죽었다'가 구분되지 않고, 판정 줄이 없으니 감사도 못 한다.
    """
    try:
        return _run_criterion(c, args)
    except Exception as e:                                    # noqa: BLE001
        return False, "criterion error (%s: %s) check=%r" % (
            type(e).__name__, str(e)[:160], c.get("check", ""))


def _run_criterion(c, args):
    kind = c.get("check", "")
    if kind == "files":
        g = resolve(c.get("glob", ""), args)
        need, err = _int_field(c, args, "min", 1)
        if err:
            return False, err
        minb, err = _int_field(c, args, "minbytes", 1)
        if err:
            return False, err
        files = [f for f in glob.glob(g) if os.path.isfile(f) and os.path.getsize(f) >= minb]
        return len(files) >= need, "files matched=%d (need>=%d, >=%dB)" % (len(files), need, minb)
    if kind == "contains":
        f = resolve(c.get("file", ""), args)
        # pat도 자리표시자를 치환한다 — file만 치환하면 `커버리지 [0-9]+/{roster}`처럼
        # **워커가 낸 값과 대조하는 기준**을 쓸 수가 없다(리터럴로 남아 영원히 0건이 되고,
        # 그러면 "안 맞아서 막혔다"와 "기준이 죽어 있다"가 구분되지 않는다).
        # 치환값은 기본적으로 `re.escape` 된다 — 이유는 `resolve_re` 참조.
        pat = resolve_re(c.get("pat", ""), args, raw=bool(c.get("raw_pat")))
        need, err = _int_field(c, args, "min", 1)
        if err:
            return False, err
        if not f or not os.path.isfile(f):
            return False, "file not found: %s" % f
        try:
            with open(f, encoding="utf-8", errors="ignore") as fh:
                txt = fh.read()
        except OSError as e:
            return False, "read error: %s" % e
        n = len(re.findall(pat, txt))
        return n >= need, "pattern hits=%d (need>=%d) file=%s" % (n, need, f)
    if kind == "exists":
        p = resolve(c.get("path", ""), args)
        if p and os.path.isdir(p):
            ok = len(os.listdir(p)) > 0
        elif p:
            ok = os.path.isfile(p) and os.path.getsize(p) > 0
        else:
            ok = False
        return ok, "exists+nonempty=%s path=%s" % (ok, p)
    if kind == "xlsx_links":
        # xlsx 출처 규율의 *결정론* floor: 하이퍼링크 수·distinct 타깃·셀 라벨 패턴(예 R\d+) 카운트.
        # 'R# 인덱스 없는 맨 링크/무출처' 같은 형식 미준수를 LLM 없이 직접 차단(openpyxl).
        f = resolve(c.get("file", ""), args)
        need, err = _int_field(c, args, "min", 1)
        if err:
            return False, err
        need_d, err = _int_field(c, args, "min_distinct", 0)
        if err:
            return False, err
        # cell_pat도 치환한다 — 안 하면 `R{roster}` 같은 기준이 리터럴로 남아 영원히 0건이다.
        cell_pat = resolve_re(c.get("cell_pat"), args, raw=bool(c.get("raw_pat"))) \
            if c.get("cell_pat") else None
        cell_min, err = _int_field(c, args, "cell_min", 0)
        if err:
            return False, err
        if not f or not os.path.isfile(f):
            return False, "file not found: %s" % f
        try:
            import openpyxl
            wb = openpyxl.load_workbook(f)
        except Exception as e:
            return False, "xlsx 로드 실패(openpyxl): %s" % e
        links = []; cellhits = 0
        rx = re.compile(cell_pat) if cell_pat else None
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.hyperlink and getattr(cell.hyperlink, "target", None):
                        links.append(cell.hyperlink.target)
                    if rx is not None and cell.value is not None and rx.search(str(cell.value)):
                        cellhits += 1
        distinct = len(set(links))
        ok = (len(links) >= need) and (distinct >= need_d) and (rx is None or cellhits >= cell_min)
        return ok, "xlsx hyperlinks=%d distinct=%d cellpat<%s>hits=%d (need link>=%d distinct>=%d cell>=%d)" % (
            len(links), distinct, (cell_pat or "-"), cellhits, need, need_d, cell_min)
    if kind == "xlsx_highlight":
        # 하이라이트(HIGHLIGHT fill) 과다의 *결정론* floor: 한 시트에서 *연속된* 하이라이트 행이
        # max_run을 넘으면 FAIL. "표 절반 칠하기/의미 없는 하이라이트"(color-policy 위반)를 LLM 없이 차단.
        # 떨어진 표가 각자 1~2행 하이라이트면 사이 헤더·빈행이 run을 끊어 통과(정상). 3행+ 연속 = 과다.
        f = resolve(c.get("file", ""), args)
        rgb6 = str(c.get("rgb", "D9F2D0"))[-6:].upper()
        max_run, err = _int_field(c, args, "max_run", 2)
        if err:
            return False, err
        if not f or not os.path.isfile(f):
            return False, "file not found: %s" % f
        try:
            import openpyxl
            wb = openpyxl.load_workbook(f)
        except Exception as e:
            return False, "xlsx 로드 실패(openpyxl): %s" % e
        worst = 0; worst_loc = ""
        for ws in wb.worksheets:
            run = 0; run_start = None
            for row in ws.iter_rows():
                row_has = False; rownum = None
                for cell in row:
                    rownum = cell.row
                    fl = cell.fill
                    if fl is not None and getattr(fl, "patternType", None) == "solid":
                        rgb = getattr(fl.start_color, "rgb", None)
                        if isinstance(rgb, str) and rgb[-6:].upper() == rgb6:
                            row_has = True
                            break
                if row_has:
                    if run == 0:
                        run_start = rownum
                    run += 1
                    if run > worst:
                        worst = run; worst_loc = "%s!rows%d-%d" % (ws.title, run_start, rownum)
                else:
                    run = 0
        ok = worst <= max_run
        return ok, "max consecutive highlight run=%d (allow<=%d) at %s" % (worst, max_run, worst_loc or "-")
    if kind == "xlsx_no_crammed_src":
        # 출처 규율의 *결정론* floor: 복수 출처를 한 셀에 몰아넣으면 FAIL — 출처가 여럿이면
        # 출처1/출처2 *열*로 분리(styles.make_table link_cols)하거나 하위 행으로. sources_crammed
        # 자기신고(attest) 오신고를 LLM 없이 직접 차단. 두 형태 모두 탐지:
        #   (a) 한 대괄호 내 2+ 출처([R16·R84]).
        #   (b) 셀 전체가 R코드 나열인 순수 출처태그(대괄호 없이 'R1·R2·R3'·'R1, R2'·'R1/R2' 등,
        #       구분자 무관)인데 그 셀 하이퍼링크 수 < R코드 수(각 코드가 개별 클릭 안 됨).
        # 별개 대괄호 병기([R37] … [R37c])·서술문 안 인라인 인용(…[R12]…[R25]…)은 정상(미검출).
        f = resolve(c.get("file", ""), args)
        if not f or not os.path.isfile(f):
            return False, "file not found: %s" % f
        try:
            import openpyxl
            wb = openpyxl.load_workbook(f)
        except Exception as e:
            return False, "xlsx 로드 실패(openpyxl): %s" % e
        rx = re.compile(r"\[R\d+[a-z]?(?:\s*[·,]\s*R\d+[a-z]?)+\]")          # (a) 대괄호 크램
        rx_code = re.compile(r"R\d+[a-z]?")
        rx_taglist = re.compile(r"R\d+[a-z]?(?:\s*[·,/、]\s*R\d+[a-z]?)+")   # (b) 순수 출처태그 나열
        hits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    s = str(cell.value)
                    crammed = bool(rx.search(s))                            # (a)
                    if not crammed:
                        stripped = s.replace("[", "").replace("]", "").strip()
                        if rx_taglist.fullmatch(stripped):                  # 셀 전체가 코드 나열
                            n_codes = len(rx_code.findall(stripped))
                            n_links = 1 if cell.hyperlink else 0            # openpyxl: 셀당 링크 최대 1
                            if n_codes >= 2 and n_links < n_codes:          # (b) 링크<코드
                                crammed = True
                    if crammed:
                        hits.append("%s!%s" % (ws.title, cell.coordinate))
        ok = len(hits) == 0
        return ok, "crammed 복수출처 셀(대괄호/가운뎃점·공백 등 한 셀 몰림·링크<코드)=%d (0이어야 함; 출처1/출처2 열 분리) %s" % (
            len(hits), (", ".join(hits[:6]) if hits else ""))
    if kind == "deliverable_name":
        # 산출물 버전 규율의 *결정론* floor: deliverable 파일명이 vN_M 형식인지 검사.
        # flat _vN(v2)·base(무버전)·_wip은 FAIL — 덮어쓰기 우회/버전 규율 미준수를 *판정 시점*에 하드 차단.
        # 워커가 빌드 코드로 styles.save_deliverable 가드를 우회해도 이 게이트가 사인오프를 막는다.
        f = resolve(c.get("file", "{deliverable}"), args)
        pat = c.get("pat", r"_v\d+_\d+\.[A-Za-z0-9]+$")
        base = os.path.basename(f) if f else ""
        # ★ 파일명만 보고 **존재는 안 봤다.** 그러면 있지도 않은 경로를 그럴듯한 이름으로
        #   내면 통과한다 — 산출물 버전 규율을 검사하는 자리인데 산출물이 없어도 된다는 뜻이었다.
        if not f or not os.path.isfile(f):
            return False, "deliverable 파일이 없다: %s" % (f or "(경로 없음)")
        if os.path.getsize(f) <= 0:
            return False, "deliverable이 빈 파일이다: %s" % f
        ok = bool(re.search(pat, base))
        return ok, "deliverable 파일명 vN_M 형식=%s (pat<%s> name=%s)" % (ok, pat, base or "-")
    return False, "unknown check '%s'" % kind


# 순서 토큰의 수명. 한 워크플로가 이보다 오래 걸리면 앞 단계를 다시 통과해야 한다.
TOKEN_TTL_SEC = 12 * 3600


def _token_sig(tok):
    body = json.dumps({k: tok.get(k) for k in ("skill", "cp", "verdict", "wf", "epoch")},
                      sort_keys=True, ensure_ascii=False)
    return hmac.new(_form_key(), body.encode("utf-8"), hashlib.sha256).hexdigest()


def read_token(skill, wf=""):
    """순서 토큰. **서명·만료·귀속이 어긋나면 없는 것으로 취급한다.**

    ★ 예전에는 서명도 만료도 없었다. 그래서 ① 손으로 `{"cp":"dispatch","verdict":"PASS"}`를
    써 넣으면 순서 강제가 통째로 무력화되고 ② 며칠 전 토큰이 그대로 유효했다
    (실측: 9월 3일자 `skill-maintainer` 토큰이 9월 10일에도 살아 있었다).
    순서 강제는 "직전 단계를 *이번 작업에서* 통과했다"를 뜻해야 의미가 있다.
    """
    p = ledger_path(skill, wf)
    if not os.path.exists(p):
        return None
    try:
        with open(p, encoding="utf-8") as fh:
            line = fh.read().strip()
        tok = json.loads(line) if line else None
    except (ValueError, OSError):
        return None
    if not isinstance(tok, dict):
        return None
    if "sig" not in tok:
        tok["_invalid"] = "서명 없음(손으로 쓴 토큰이거나 구버전)"
        return tok
    if not hmac.compare_digest(str(tok.get("sig")), _token_sig(tok)):
        tok["_invalid"] = "서명 불일치(내용이 바뀌었다)"
        return tok
    if (tok.get("skill"), tok.get("wf")) != (skill, wf):
        tok["_invalid"] = "다른 스킬·워크플로의 토큰"
        return tok
    age = time.time() - float(tok.get("epoch") or 0)
    if age > TOKEN_TTL_SEC:
        tok["_invalid"] = "만료(%.1f시간 전 발급)" % (age / 3600)
        return tok
    return tok


def write_token(skill, cp, wf=""):
    tok = {"skill": skill, "wf": wf, "cp": str(cp), "verdict": "PASS",
           "epoch": int(time.time()), "ts": time.strftime("%Y-%m-%d %H:%M:%S")}
    tok["sig"] = _token_sig(tok)
    with open(ledger_path(skill, wf), "w", encoding="utf-8") as fh:
        fh.write(json.dumps(tok, ensure_ascii=False) + "\n")


def remove_ledger(skill, wf=""):
    try:
        os.remove(ledger_path(skill, wf))
    except OSError:
        pass


def render_prompt(tmpl, args):
    """judge_prompt의 {id} 자리표시자를 워커 제출 경로/값으로 치환."""
    out = tmpl
    for k, v in args.items():
        out = out.replace("{%s}" % k, str(v))
    return out


CP_KEYS = {"after", "criteria", "intent", "mode", "observations"}
OBS_KEYS = {"ask", "attest_must", "check", "file", "id", "minlen", "variants"}
VAR_KEYS = {"ask", "must"}
CRIT_KEYS = {"cell_min", "cell_pat", "check", "file", "glob", "max_run", "min",
             "min_distinct", "minbytes", "pat", "path", "raw_pat"}
CHECK_KINDS = {"contains", "deliverable_name", "exists", "files", "xlsx_highlight",
               "xlsx_links", "xlsx_no_crammed_src"}


def validate_cp(skill, cp, cpr):
    """루브릭 구조를 로드 시점에 검사한다. 어긋나면 **GATE ERROR로 세운다.**

    막는 것 둘.
    ① **오타로 인한 조용한 강등** — `variants`를 `varients`로 쓰면 그 항목이 채점 대상에서
       빠져 `inputs`로 간다. 그런데 **질문은 그대로 출제되므로** 워커·감사자 모두
       "물었고 답했다"만 보고 채점되지 않은 사실을 알 수 없다.
    ② **아무것도 검사하지 않는 체크포인트** — `criteria`도 채점 항목도 없으면 인자 0개로
       PASS가 난다. 검사가 없는 게이트는 통과가 아니라 **정의 오류**다.
    """
    errs = []
    bad = set(cpr.keys()) - CP_KEYS
    if bad:
        errs.append("cp에 모르는 키 %s (오타? 가능: %s)" % (sorted(bad), sorted(CP_KEYS)))
    obs = cpr.get("observations") or []
    ids = []
    for i, o in enumerate(obs):
        if not isinstance(o, dict):
            errs.append("observations[%d]가 객체가 아니다" % i)
            continue
        if not o.get("id"):
            errs.append("observations[%d]에 id가 없다" % i)
        ids.append(o.get("id"))
        ob = set(o.keys()) - OBS_KEYS
        if ob:
            errs.append("observation '%s'에 모르는 키 %s — 오타면 그 항목이 "
                        "**질문은 되고 채점은 안 된다**" % (o.get("id"), sorted(ob)))
        for j, v in enumerate(o.get("variants") or []):
            if not isinstance(v, dict):
                errs.append("'%s'.variants[%d]가 객체가 아니다" % (o.get("id"), j))
                continue
            vb = set(v.keys()) - VAR_KEYS
            if vb or not v.get("ask") or not v.get("must"):
                errs.append("'%s'.variants[%d]: ask·must가 둘 다 있어야 하고 "
                            "다른 키는 못 온다(현재 %s)" % (o.get("id"), j, sorted(v.keys())))
    dup = {i for i in ids if i and ids.count(i) > 1}
    if dup:
        errs.append("observation id 중복 %s — 뒤가 앞을 덮어 하나가 채점되지 않는다" % sorted(dup))
    crits = cpr.get("criteria") or []
    for i, c in enumerate(crits):
        if not isinstance(c, dict):
            errs.append("criteria[%d]가 객체가 아니다" % i)
            continue
        cb = set(c.keys()) - CRIT_KEYS
        if cb:
            errs.append("criteria[%d]에 모르는 키 %s" % (i, sorted(cb)))
        if c.get("check") not in CHECK_KINDS:
            errs.append("criteria[%d]: 모르는 check '%s' — 이 기준은 항상 FAIL이 된다"
                        % (i, c.get("check")))
    _in, grad = _split_obs([o for o in obs if isinstance(o, dict)])
    if not crits and not grad:
        errs.append("검사가 하나도 없다(criteria 0 · 채점 항목 0) — "
                    "인자 없이 PASS가 나므로 게이트가 아니다")
    if errs:
        print("GATE ERROR [%s:%s] 루브릭 구조 오류 %d건:" % (skill, cp, len(errs)))
        for e in errs:
            print("  · %s" % e)
        print(">>> 루브릭을 고쳐라. 구조가 깨진 게이트는 통과시키지 않는다.")
        sys.exit(3)


def get_cp(skill, cp):
    r = load_rubric(skill)
    if r is None:
        print("GATE ERROR: 루브릭 없음 — _stepgate/rubrics/%s.json" % skill)
        sys.exit(3)
    cpr = r.get(cp)
    if not cpr:
        print("GATE ERROR: 체크포인트 '%s' 없음 (가능: %s)" % (cp, ", ".join(r.keys())))
        sys.exit(3)
    validate_cp(skill, cp, cpr)
    return cpr


# ---- 비동기 감사(T2/T3) 인프라 — 큐·wfid·audit_rubric ---------------------

AUDIT_RUBRIC_DIR = os.path.normpath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "audit_rubrics"))


def _safe(s):
    return "".join(c if (c.isalnum() or c in "-_.") else "_" for c in str(s))


def audit_queue_dir():
    d = os.path.join(gate_dir(), "audit_queue")
    os.makedirs(d, exist_ok=True)
    return d


def deliv_hash(path):
    norm = os.path.normpath(os.path.abspath(path)).lower()
    return hashlib.sha1(norm.encode("utf-8", "replace")).hexdigest()[:10]


def _series_stem(path):
    """산출물 '계열' 키 = 파일명에서 버전 표기(_vN_M)만 제거(확장자 유지·소문자).
    같은 산출물의 서로 다른 버전(v3_1·v3_2…)을 한 계열로 묶어, enqueue 시 *이전 버전* 감사
    엔트리를 SUPERSEDE하기 위함(새 파일명으로 버전을 올려도 구버전 좀비가 큐에 안 남게).
    버전 표기가 없으면 파일명 그대로 → 계열=자기 자신(기존 deliv_hash와 동일 동작, 오탐 없음)."""
    b = os.path.basename(path or "").strip().lower()
    m = re.search(r"_v\d+_\d+(\.[a-z0-9]+)$", b)
    if m:
        return b[:m.start()] + m.group(1)
    return b


def wfid_file(skill):
    return os.path.join(gate_dir(), _safe(skill) + ".wfid")


def _wfid(skill):
    """워크플로 파티션 키(= 큐 파일명). 세션 sid를 못 얻는 환경의 안정 폴백 — init/clear로 갱신."""
    p = wfid_file(skill)
    if os.path.isfile(p):
        try:
            v = open(p, encoding="utf-8").read().strip()
            if v:
                return v
        except OSError:
            pass
    wid = "wf_%s_%d" % (_safe(skill), int(time.time() * 1000))
    try:
        with open(p, "w", encoding="utf-8") as fh:
            fh.write(wid)
    except OSError:
        pass
    return wid


def wfid_queue_path(skill):
    return os.path.join(audit_queue_dir(), _safe(_wfid(skill)) + ".jsonl")


def queue_path(skill, wf):
    """감사 큐 = **skill + 워크플로 키(wf)** 별 파일(ledger/form과 동일 파티션).
    ★ skill을 키에 포함해야 같은 과제폴더(dir)에서 서로 다른 스킬(예: research-verified·
    research-docx)이 동시 진행해도 큐 파일을 공유·삭제하지 않는다 — init/clear의
    _clear_workflow가 *공유* 큐 파일을 통째로 os.remove 하며 타 스킬의 PENDING 감사
    엔트리까지 지우던 교차오염 버그 수정. reqid는 skill을 임베드하므로 audit-judge의
    전수 스캔(_find_queue_entry)은 파일이 갈려도 그대로 찾는다."""
    return os.path.join(audit_queue_dir(), _safe(skill + "__" + wf) + ".jsonl")


def _wf_key(skill, args):
    """워크플로 파티션 키(동시 워크플로 격리). 우선순위:
    명시 wf= > dir=(과제폴더) 해시 > deliverable 부모폴더 해시 > .wfid 폴백(스킬당 1개·단일작업 호환).
    같은 과제폴더의 동시 작업은 같은 키(한 워크플로로 간주) — 산출물 구분은 deliv_hash로."""
    wf = (args.get("wf") or "").strip()
    if wf:
        return "wf_" + _safe(wf)[:32]
    d = (args.get("dir") or "").strip()
    if not d:
        deliv = (args.get("deliverable") or "").strip()
        if deliv:
            d = os.path.dirname(deliv)
    if d:
        return "d_" + deliv_hash(d)
    return _wfid(skill)


def append_queue(path, obj):
    with open(path, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(obj, ensure_ascii=False) + "\n")


def read_queue_collapsed(path):
    """append-only 큐를 reqid별 최신 라인으로 접어 dict 반환(.ledger와 동일 패턴)."""
    out = {}
    if not os.path.isfile(path):
        return out
    try:
        with open(path, encoding="utf-8", errors="ignore") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                try:
                    o = json.loads(line)
                except ValueError:
                    continue
                rid = o.get("reqid")
                if rid:
                    out[rid] = o
    except OSError:
        pass
    return out


def load_audit_rubric(skill):
    p = os.path.join(AUDIT_RUBRIC_DIR, skill + ".json")
    if not os.path.isfile(p):
        return None
    try:
        with open(p, encoding="utf-8") as fh:
            return json.load(fh)
    except (ValueError, OSError) as e:
        print("AUDIT ERROR: audit_rubric 로드 실패 %s :: %s" % (p, e))
        return None


# 봉인 정답지를 가리키는 문자열인가. `audit_rubrics/`는 비봉인이라 제외한다.
_RX_RUBRIC_BAK = re.compile(r"[_\-.]rubric[s]?[_\-.].*backup|backup.*[_\-.]rubric[s]?[_\-.]", re.I)
# 셸로 파일을 읽는 흔한 수단. 이 목록에 없는 방법이 남는 것은 인정한다(탐지지 예방 아님).
_RX_SHELL_READ = re.compile(
    r"\b(cat|bat|head|tail|less|more|nl|strings|grep|rg|ag|awk|sed|jq|xxd|od|type)\b"
    r"|python[0-9.]*\s+-c|open\s*\(|read_text|readlines|json\.load", re.I)


def _refs_sealed_rubric(text):
    """봉인 루브릭(또는 그 **백업 사본**)을 가리키는가.

    ★ 두 구멍을 막는다. ① 경로가 `_stepgate/rubrics`가 아니어도 파일명이
    `_<skill>_rubric_backup_*.json`이면 그것은 **정답지 전체 사본**이다 —
    다른 폴더에 복사해 두고 읽으면 R2가 못 봤다. ② `audit_rubrics/`는 비봉인이므로
    계속 제외한다(감사자가 읽는 것이 정상이다).
    """
    if not text:
        return False
    t = text.replace("\\", "/")
    if "audit_rubrics" in t:
        return False
    if "_stepgate" in t and "rubrics" in t:
        return True
    return bool(_RX_RUBRIC_BAK.search(t))


def _reads_via_shell(text):
    """셸 명령이 파일을 *읽는* 형태인가(쓰기·이동만이면 아니다)."""
    return bool(_RX_SHELL_READ.search(text or ""))


def _find_queue_entry(reqid):
    """reqid가 든 큐 파일을 전수 스캔으로 찾아 (path, obj) 반환."""
    for p in glob.glob(os.path.join(audit_queue_dir(), "*.jsonl")):
        coll = read_queue_collapsed(p)
        if reqid in coll:
            return p, coll[reqid]
    return None, None


def _reqid_parts(reqid):
    """신형 reqid(`skill__cp__wf__ts`)에서 (skill, cp, wf) 복원.
    큐 파일이 유실돼도 audit-judge가 기록을 이어가게 하는 복구용. 구형(`_` 구분·wf 절단)은
    복원 불가 → None(그 경우 기존처럼 에러). wf가 `__`를 포함할 여지까지 고려해 가운데를 모두 wf로."""
    parts = str(reqid).split("__")
    if len(parts) < 4:
        return None
    skill, cp, wf = parts[0], parts[1], "__".join(parts[2:-1])
    if not (skill and cp and wf):
        return None
    return skill, cp, wf


def _clear_workflow(skill, wf=""):
    """이 워크플로의 감사 큐 제거. wf 지정 시 *그 워크플로 큐만*(동시 진행 중인 타 워크플로 큐는 보존),
    미지정이면 폴백 .wfid 큐 + .wfid 파일을 제거(단일작업 종료)."""
    qp = queue_path(skill, wf) if wf else wfid_queue_path(skill)
    try:
        os.remove(qp)
    except OSError:
        pass
    if not wf:
        try:
            os.remove(wfid_file(skill))
        except OSError:
            pass


def _enqueue(skill, cp, args):
    """PENDING 큐 라인 1줄 기록 → reqid 반환. 큐 파일은 워크플로 키(wf)별로 분리(동시작업 격리).
    *같은 cp*의 이전 활성 엔트리 중 **같은 산출물(deliv_hash) 또는 같은 계열(파일명 _vN_M만 다른
    버전)**을 SUPERSEDED 처리(재빌드·버전상승 시 옛 FAIL/PENDING이 RESOLVE를 막지 않게) — 다른
    산출물 계열은 공존(서로 죽이지 않음). 계열 SUPERSEDE가 없으면 v3_1→v3_2처럼 파일명이 바뀔 때
    구버전 감사가 좀비로 누적돼 폴더 resolve를 계속 막던 문제를 근본 차단. deliverable= 필수."""
    deliverable = args.get("deliverable", "").strip()
    wf = _wf_key(skill, args)
    qp = queue_path(skill, wf)
    dh = deliv_hash(deliverable) if deliverable else ""
    ss = _series_stem(deliverable) if deliverable else ""
    coll = read_queue_collapsed(qp)
    for rid, o in coll.items():
        if o.get("cp") != cp or o.get("status") == "SUPERSEDED":
            continue
        same_deliv = bool(dh) and o.get("deliv_hash") == dh
        o_ss = _series_stem(o.get("deliverable", ""))
        same_series = bool(ss) and bool(o_ss) and o_ss == ss
        if same_deliv or same_series:
            sup = dict(o)
            sup["status"] = "SUPERSEDED"
            sup["ts_done"] = time.strftime("%Y-%m-%d %H:%M:%S")
            append_queue(qp, sup)
    # reqid에 *전체 wf*를 `__`로 임베드 — 배경 감사 도중 부모가 clear/init/재빌드로 큐 파일을
    # 지워도(같은 스킬 워크플로) audit-judge가 reqid만으로 엔트리를 복구해 결과를 기록하게(유실 방지).
    reqid = "%s__%s__%s__%d" % (skill, cp, wf, int(time.time() * 1000))
    obj = {"reqid": reqid, "skill": skill, "cp": cp,
           "deliverable": deliverable, "deliv_hash": dh,
           "wfid": wf, "dir": args.get("dir", ""),
           "status": "PENDING", "subverdict": None, "findings": "", "evidence": "",
           "ts_enq": time.strftime("%Y-%m-%d %H:%M:%S"), "ts_done": None, "agent": None}
    append_queue(qp, obj)
    return reqid


def _print_fire_block(skill, cp, reqid, deliverable):
    print("GATE ENQUEUED [%s:%s] req=%s → 비동기 빡빡 감사를 *백그라운드 에이전트*로 지금 발사하고 "
          "다음 작업을 시작하라(대기 금지):" % (skill, cp, reqid))
    print("----- FIRE (Agent run_in_background=true, subagent_type='general-purpose') -----")
    print("[주의] gate-audit은 Skill이라 subagent_type이 아니다 → general-purpose(또는 claude) 백그라운드 에이전트에 "
          "이 지시를 그대로 위임하라(아래 패킷에 검사표가 다 들어 있어 동일하게 검수된다).")
    print("너는 백그라운드 감사자다. 산출물=%s 를 audit_rubric의 모든 lens로 *적대적*으로 검수하라." % deliverable)
    print("1) `SG audit-judge %s` 로 감사 패킷(lens·directives)을 받는다." % reqid)
    print("2) 산출물을 실제로 열고 lens별로 위반을 *그대로 인용*(라인/셀좌표/캡션). 인용 못 하면 그 lens는 PASS.")
    print("3) needs_fetch lens는 링크를 실제 WebFetch해 수치 착지 확인.")
    print("4) 종료: `SG audit-judge %s --record subverdict=PASS|FAIL findings=<위반요약> evidence=<핵심 인용>`" % reqid)
    print("----- END FIRE -----")


# ---- commands -------------------------------------------------------------

def form_state_path(skill, cp, wf=""):
    key = skill + "__" + (wf + "__" if wf else "") + cp
    safe = "".join(c if (c.isalnum() or c in "-_.") else "_" for c in key)
    return os.path.join(gate_dir(), safe + ".form")


# 출제 상태(.form)의 수명. 이 안에 judge를 내야 한다.
FORM_TTL_SEC = 3 * 3600


def _form_key():
    """출제 상태 서명 키. 없으면 만든다(0600).

    **정직한 한계**: 워커와 같은 파일시스템에 있으므로 암호벽이 아니다(봉인 루브릭과 같은
    전제 — 협조적·부주의 모델). 이것이 막는 것은 ① **손으로 쓴 출제 파일**과
    ② 다른 체크포인트·워크플로의 출제를 갖다 붙이는 것, ③ **낡은 출제의 재사용**이다.
    극성 회전은 "이번 회차에 무엇을 물었나"가 정확할 때만 의미가 있고, 그 기록을
    아무나 쓸 수 있으면 회전 자체가 장식이 된다.
    """
    p = os.path.join(gate_dir(), ".formkey")
    if not os.path.exists(p):
        fd = os.open(p, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(binascii.hexlify(os.urandom(32)).decode())
    with open(p, encoding="utf-8") as fh:
        return fh.read().strip().encode()


def _form_sig(payload):
    """서명 대상은 '무엇을·어느 극성으로·언제 물었나' 전부다."""
    body = json.dumps({k: payload[k] for k in ("skill", "cp", "wf", "asked", "variant", "ts")},
                      sort_keys=True, ensure_ascii=False)
    return hmac.new(_form_key(), body.encode("utf-8"), hashlib.sha256).hexdigest()


def write_form_state(skill, cp, wf, asked, variant):
    st = {"skill": skill, "cp": cp, "wf": wf, "asked": asked,
          "variant": variant, "ts": int(time.time())}
    st["sig"] = _form_sig(st)
    with open(form_state_path(skill, cp, wf), "w", encoding="utf-8") as fh:
        json.dump(st, fh, ensure_ascii=False)


def read_form_state(skill, cp, wf):
    """(상태, 오류). 서명·만료·귀속이 하나라도 어긋나면 상태를 **주지 않는다.**"""
    p = form_state_path(skill, cp, wf)
    if not os.path.exists(p):
        return None, None
    try:
        with open(p, encoding="utf-8") as fh:
            st = json.load(fh)
    except (ValueError, OSError) as e:
        return None, "출제 기록을 읽을 수 없다(%s)" % e
    if not isinstance(st, dict) or "sig" not in st:
        return None, "출제 기록에 서명이 없다 — 손으로 쓴 파일이거나 구버전이다"
    try:
        expect = _form_sig(st)
    except KeyError as e:
        return None, "출제 기록에 필드가 빠졌다(%s)" % e
    if not hmac.compare_digest(str(st.get("sig")), expect):
        return None, "출제 기록의 서명이 맞지 않다 — 내용이 바뀌었다"
    if (st.get("skill"), st.get("cp"), st.get("wf")) != (skill, cp, wf):
        return None, "다른 체크포인트·워크플로의 출제 기록이다"
    age = time.time() - int(st.get("ts") or 0)
    if age > FORM_TTL_SEC:
        return None, "출제가 만료됐다(%.1f시간 전) — 낡은 출제로는 채점하지 않는다" % (age / 3600)
    return st, None


# 사후 감사가 '산출물'로 셀 확장자. **`.md`가 빠져 있던 것이 큰 구멍이었다** —
# R1(게이트 미호출)·R3(FAIL 무시)·R6(비동기감사 미완료)·R7(해소게이트 누락)이 전부
# `deliverables`가 비면 발동하지 않으므로, 산출물이 `.md`뿐인 스킬(trade-run·world-study·
# fact-check)에서는 **7규칙 중 4개가 죽어 있었다.**
DELIVERABLE_EXT = r"\.(xlsx|xlsm|docx|pptx|md)$"
# 산출물이 아닌 `.md` — 제외하지 않으면 스크래치·백업·리포 메타 파일이 오탐을 만든다.
_NOT_DELIVERABLE_BASENAMES = {"claude.md", "readme.md", "memory.md", "changelog.md"}
_NOT_DELIVERABLE_PARTS = ("scratchpad", "_raw_sources", "claude_stepgate",
                          os.sep + ".claude" + os.sep, "/memory/", os.sep + "memory" + os.sep)


def is_deliverable_path(fp):
    """이 경로가 사후 감사에서 '산출물'로 셀 대상인가.

    판정을 한 곳에 모으는 이유: 예전에는 정규식 한 줄이 `xlsx|docx|pptx`만 잡아서
    `.md` 스킬의 감사가 조용히 꺼져 있었다. 확장자를 넓히면 이번엔 **스크래치·백업이
    산출물로 잡혀 오탐**이 나므로 제외 규칙도 같은 자리에 적어야 한다.
    """
    if not fp or not re.search(DELIVERABLE_EXT, fp, re.I):
        return False
    low = fp.replace("\\", os.sep).lower()
    base = os.path.basename(low)
    if base in _NOT_DELIVERABLE_BASENAMES:
        return False
    if base.startswith("_"):                 # `_*_backup_*`·`_tmp`·작업 파일
        return False
    if any(part.lower() in low for part in _NOT_DELIVERABLE_PARTS):
        return False
    if low.startswith("/tmp/") or low.startswith("/private/tmp/") or "/var/folders/" in low:
        return False
    return True


def quote_check(file_path, snippet, minlen):
    """워커가 낸 *문구*가 산출물 파일에 literal로 있는가(어감·암기로 못 때움·매 회차 다름)."""
    snip = (snippet or "").strip()
    if not file_path or not os.path.isfile(file_path):
        return False, "file not found: %s" % file_path
    if len(snip) < minlen:
        return False, "인용이 너무 짧음(>=%d자 필요)" % minlen
    # xlsx/xlsm은 바이너리(zip) — 셀 값·하이퍼링크 타깃을 텍스트 blob으로 추출해 대조(openpyxl).
    if file_path.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            parts.append(str(cell.value))
                        if cell.hyperlink and getattr(cell.hyperlink, "target", None):
                            parts.append(str(cell.hyperlink.target))
            txt = "\n".join(parts)
        except Exception as e:
            return False, "xlsx 인용 검사 실패(openpyxl): %s" % e
    else:
        try:
            with open(file_path, encoding="utf-8", errors="ignore") as fh:
                txt = fh.read()
        except OSError as e:
            return False, "read error: %s" % e
    ok = snip in txt
    return ok, ("인용이 파일에 존재" if ok else "인용이 파일에 없음(허위/오타?)")


def _split_obs(obs):
    """observations → (inputs: 채점 안 하는 값/경로, gradeables: attest/variants/quote 채점 대상)."""
    inputs, gradeables = [], []
    for o in obs:
        if "attest_must" in o or "variants" in o or o.get("check") == "quote":
            gradeables.append(o)
        else:
            inputs.append(o)
    return inputs, gradeables


def _variants(o):
    """관찰 항목의 극성 변형 목록(없으면 빈 리스트 → 레거시 ask+attest_must 경로)."""
    v = o.get("variants")
    return v if isinstance(v, list) and v else []


def _has_variants(cpr):
    """이 체크포인트에 극성 회전 항목이 하나라도 있는가 → form 선행 강제 여부."""
    return any(_variants(o) for o in cpr.get("observations", []))


def _log_wipe(action, skill, wf):
    """`init`/`clear`가 **무엇을 지웠는지** 지워지지 않는 곳에 남긴다.

    ★ 예전에는 둘 다 **인증 없는 FAIL 지우기 버튼**이었다. 기록된 FAIL을 만난 워커가
    `clear`를 부르고 처음부터 다시 하면 아무 흔적도 남지 않았다. 사후 감사가
    "FAIL을 무시했다"(R3)를 탐지하려면 **지운 사실 자체가 증거로 남아야** 한다.
    이 로그는 `clear`가 지우지 않는다(그러면 같은 구멍이 된다).
    """
    tok = read_token(skill, wf)
    pend = 0
    try:
        for _rid, o in read_queue_collapsed(queue_path(skill, wf)).items():
            if o.get("status") == "PENDING":
                pend += 1
    except Exception:                                       # noqa: BLE001
        pass
    row = {"ts": time.strftime("%Y-%m-%d %H:%M:%S"), "action": action,
           "skill": skill, "wf": wf,
           "token_cp": (tok or {}).get("cp"), "token_invalid": (tok or {}).get("_invalid"),
           "pending_audits": pend}
    try:
        with open(os.path.join(gate_dir(), "wipe_log.jsonl"), "a", encoding="utf-8") as fh:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")
    except OSError:
        pass
    if (tok and not tok.get("_invalid")) or pend:
        print("  ※ 지운 것: 토큰 cp=%s · PENDING 감사 %d건 — wipe_log.jsonl에 기록했다."
              % ((tok or {}).get("cp"), pend))


def cmd_init(skill, kv_args):
    args = parse_kv(kv_args)
    wf = _wf_key(skill, args)
    _log_wipe("init", skill, wf)
    remove_ledger(skill, wf)
    _clear_workflow(skill, wf)
    print("GATE INIT [%s] wf=%s — 새 워크플로(이 워크플로의 잔여 토큰·감사큐만 청소). 루브릭: %s"
          % (skill, wf, RUBRIC_DIR))


def cmd_form(skill, cp, kv_args):
    args = parse_kv(kv_args)
    wf = _wf_key(skill, args)
    cpr = get_cp(skill, cp)
    inputs, gradeables = _split_obs(cpr.get("observations", []))
    # ★ 표본 추출 없음 — 모든 관찰 항목이 매 회차 *전부* 출제된다(기준이 그 회차에 빠지는 일 없음).
    #   암기 방지는 항목을 빼는 게 아니라 *극성 회전*이 맡는다: 같은 항목을 회차마다 Y정답 표현과
    #   N정답 표현 중 하나로 물어 고정 정답 벡터를 무의미하게 만든다.
    asked = inputs + gradeables
    picked = {}                                              # id → 이번 회차에 뽑힌 변형 번호
    lines = []
    for o in asked:
        vs = _variants(o)
        if vs:
            i = random.randrange(len(vs))
            picked[o["id"]] = i
            ask = vs[i].get("ask", "")
        else:
            ask = o.get("ask", "")
        tag = " [인용: 산출물의 해당 문구를 *그대로* 붙여넣기]" if o.get("check") == "quote" else ""
        lines.append("  - %s=<...>%s  : %s" % (o["id"], tag, ask))
    # 채점 대상(gradeable) id + 뽑힌 변형 번호를 상태에 기록 → judge가 *이번 회차에 물은 극성으로* 채점
    write_form_state(skill, cp, wf, [o["id"] for o in gradeables], picked)
    print("GATE FORM [%s:%s]" % (skill, cp))
    print("취지: %s" % cpr.get("intent", ""))
    if cpr.get("after"):
        print("(직전 통과 필요: %s)" % cpr["after"])
    print("아래 항목을 정직히 답해 judge에 제출하라 — **합격 기준은 비공개**(관찰만 보고):")
    for ln in lines:
        print(ln)
    if picked:
        print("(★ 항목의 *표현이 회차마다 뒤집힌다* — 지난 회차의 답을 그대로 쓰면 틀린다. "
              "문구를 이번 회차 그대로 읽고 답하라. FAIL이면 이 출제는 폐기되고 form부터 다시다.)")
    print("제출: judge %s %s <위 인자들>" % (skill, cp))


def cmd_judge(skill, cp, kv_args):
    cpr = get_cp(skill, cp)
    args = parse_kv(kv_args)
    wf = _wf_key(skill, args)   # 동시 워크플로 격리 — 토큰·form·큐를 이 워크플로 단위로

    # (1) 순서 강제 — 직전 체크포인트 통과 토큰 확인
    after = cpr.get("after")
    if after:
        tok = read_token(skill, wf)
        bad = tok.get("_invalid") if tok else None
        if not tok or bad or str(tok.get("cp")) != str(after) or tok.get("verdict") != "PASS":
            have = (f"{tok.get('cp')} — {bad}" if bad else tok.get("cp")) if tok else "(없음)"
            print("GATE FAIL [%s:%s]: 직전 '%s' 통과 토큰 없음(현재=%s) — 순서 건너뜀?"
                  % (skill, cp, after, have))
            print(">>> 직전 체크포인트를 먼저 통과하라.")
            sys.exit(1)

    fails = []

    # (2) 디스크 재검증 (워커 도장 무시 — 진실을 직접 재도출)
    # ★ 실패한 기준의 **인덱스**를 함께 찍는다(패턴 본문은 찍지 않는다 — 정답지는 봉인이다).
    #   예전에는 `pattern hits=0 (need>=1) file=…`만 나와서 기준 40개 중 어느 것인지
    #   알 수 없었다. 2026-09-11 미국 run이 그 한 줄 때문에 멈췄고, 표기 후보를 200개
    #   넘게 추측하고도 특정하지 못했다. **인덱스만으로도 유지보수자가 지목해 고칠 수 있고,
    #   합격 기준은 여전히 안 드러난다** — 봉인은 teach-to-test를 막으려는 것이지
    #   디버깅을 막으려는 것이 아니다.
    for _i, c in enumerate(cpr.get("criteria", [])):
        ok, msg = run_criterion(c, args)
        tag = "criteria[%d]" % _i
        print("   · [disk] %s :: %s %s" % ("OK" if ok else "FAIL", tag, msg))
        if not ok:
            fails.append("%s %s" % (tag, msg))

    # (3) 채점 대상 결정 — 모든 gradeable을 채점한다(표본 추출 없음).
    #     극성 회전 항목은 form이 *이번 회차에 뽑은 변형*의 정답으로 채점 → 상태 파일이 필수다.
    _inputs, gradeables = _split_obs(cpr.get("observations", []))
    _st, _st_err = read_form_state(skill, cp, wf)
    asked_ids = set(_st.get("asked", [])) if _st else None
    picked = (_st.get("variant") or {}) if _st else {}
    if _st_err:
        # 서명·만료·귀속이 어긋난 출제 기록 — **없는 것보다 나쁘다.** 조용히 무시하면
        # 아래에서 "form을 안 돌렸다"와 구분되지 않고, 극성 회전이 장식이 된다.
        print("GATE FAIL [%s:%s]: %s → `form %s %s`부터 다시 실행하라."
              % (skill, cp, _st_err, skill, cp))
        try:
            os.remove(form_state_path(skill, cp, wf))
        except OSError:
            pass
        sys.exit(1)
    if _has_variants(cpr) and asked_ids is None:
        # form을 건너뛰고 judge를 직접 부르면 어느 극성으로 물었는지 알 수 없다 → 회전 무력화 차단.
        print("GATE FAIL [%s:%s]: 항목 표현이 회차마다 뒤집힌다 — 먼저 `form %s %s`를 실행해 이번 회차 문항을 받아라."
              % (skill, cp, skill, cp))
        sys.exit(1)
    # ★ 채점 대상이 있는데 출제 목록이 비었으면 **하드 FAIL**이다.
    #   예전에는 `asked: []`면 `graded`가 빈 리스트가 되어 **모든 각서를 조용히 건너뛰고
    #   PASS**가 났다(빈 집합은 `None`이 아니어서 위 회전 검사도 통과했다). 즉 출제 파일
    #   한 줄로 게이트 전체를 무력화할 수 있었다.
    if gradeables and asked_ids is not None and not asked_ids:
        print("GATE FAIL [%s:%s]: 채점 항목 %d개가 있는데 출제 목록이 비었다 — "
              "채점을 건너뛴 PASS는 인정하지 않는다. `form %s %s`부터 다시 실행하라."
              % (skill, cp, len(gradeables), skill, cp))
        sys.exit(1)
    graded = [o for o in gradeables if (asked_ids is None or o["id"] in asked_ids)]

    # (3a) attest(인지) = 미검증 표기 + attest_must 위반 차단 / (3b) quote = 산출물 문구 literal 대조(어감·암기 무력)
    for o in graded:
        vs = _variants(o)
        if vs:
            # 극성 회전 — 이번 회차에 뽑힌 변형의 must가 정답. 변형 번호가 없으면 채점 불가(하드 실패).
            vi = picked.get(o["id"])
            if not isinstance(vi, int) or not (0 <= vi < len(vs)):
                print("GATE FAIL [%s:%s]: '%s'의 이번 회차 출제 기록이 없다 — `form %s %s`부터 다시 실행하라."
                      % (skill, cp, o["id"], skill, cp))
                sys.exit(1)
            must = vs[vi].get("must", "")
            ans = args.get(o["id"], "")
            ok = (str(ans).strip().upper() == str(must).strip().upper())
            print("   · [attest 미검증] %s = %s (요구 %s) :: %s"
                  % (o["id"], ans or "(무응답)", must, "OK" if ok else "FAIL"))
            if not ok:
                fails.append("attest %s != %s" % (o["id"], must))
        elif "attest_must" in o:
            must = o["attest_must"]
            ans = args.get(o["id"], "")
            ok = (str(ans).strip().upper() == str(must).strip().upper())
            print("   · [attest 미검증] %s = %s (요구 %s) :: %s"
                  % (o["id"], ans or "(무응답)", must, "OK" if ok else "FAIL"))
            if not ok:
                fails.append("attest %s != %s" % (o["id"], must))
        elif o.get("check") == "quote":
            f = resolve(o.get("file", ""), args)
            ok, msg = quote_check(f, args.get(o["id"], ""), int(o.get("minlen", 12)))
            print("   · [quote 대조] %s :: %s :: %s" % (o["id"], "OK" if ok else "FAIL", msg))
            if not ok:
                fails.append("quote %s: %s" % (o["id"], msg))

    if fails:
        # ★ FAIL이면 이번 회차 출제를 폐기한다 — 답만 뒤집어 재제출하는 맹목적 시행착오 차단.
        #   재시도하려면 form을 다시 돌려야 하고, 그때 극성이 새로 굴러간다.
        try:
            os.remove(form_state_path(skill, cp, wf))
        except OSError:
            pass
        print("GATE FAIL [%s:%s] — 미달: %s" % (skill, cp, " ; ".join(fails)))
        print(">>> 다음 단계 금지. 그 단계의 실제 산출물을 만든 뒤 재시도하라(말로 done 불가).")
        print(">>> 재시도는 `form %s %s`부터 — 이번 출제는 폐기됐고 문항이 새로 나온다(답만 뒤집기 불가)."
              % (skill, cp))
        sys.exit(1)

    # (3c) async 모드 — substance를 인라인 escalate(블로킹) 대신 *비동기 빡빡 감사*로 이관.
    #   바닥(disk+attest+quote) 통과 즉시 토큰 기록 + 감사 큐 적재 → 워커는 다음 작업으로(중첩).
    # ★ async는 **루브릭만** 정한다. 워커 인자로 켤 수 있으면, 인라인 substance 심판이
    #   붙은 체크포인트를 `async=1` 한 마디로 "바닥만 보고 즉시 PASS"로 강등하고
    #   나머지를 자기가 기록하는 비동기 감사로 넘길 수 있다 — 그게 자기 채점이다.
    #   async가 필요한 스킬은 이미 루브릭 cp에 `"mode":"async"`를 갖고 있다.
    async_mode = (cpr.get("mode") == "async")
    if str(args.get("async", "")).strip() and not async_mode:
        print("GATE FAIL [%s:%s]: `async=1`로 모드를 바꿀 수 없다 — 비동기 여부는 "
              "루브릭이 정한다(이 cp는 동기다). 인자를 빼고 다시 실행하라." % (skill, cp))
        sys.exit(1)
    if async_mode:
        write_token(skill, cp, wf)
        try:
            os.remove(form_state_path(skill, cp, wf))
        except OSError:
            pass
        print("GATE PASS [%s:%s] (async floor — substance는 비동기 감사로 이관)" % (skill, cp))
        deliverable = args.get("deliverable", "").strip()
        if deliverable:
            reqid = _enqueue(skill, cp, args)
            _print_fire_block(skill, cp, reqid, deliverable)
        else:
            print(">>> [경고] deliverable= 미제출 → 비동기 감사 큐 적재 생략. "
                  "발행 직전 RESOLVE가 0건으로 FAIL 처리된다(deliverable=<abspath> 주고 재실행 권장).")
        return

    # (4) escalation — 형식 통과 후 substance 심판 (레거시 — 루브릭 mode≠async일 때만)
    if cpr.get("escalate"):
        sv = args.get("subverdict", "").strip().upper()
        if sv not in ("PASS", "FAIL"):
            print("GATE NEEDS-JUDGE [%s:%s] — 형식검사 통과. 이제 *substance* 심판이 필요하다." % (skill, cp))
            print("아래 프롬프트로 judge 서브에이전트(별도 컨텍스트·산출물 실물 읽기)를 띄워 판정시켜라:")
            print("----- JUDGE PROMPT -----")
            print(render_prompt(cpr.get("judge_prompt", "(judge_prompt 미정의)"), args))
            print("----- END JUDGE PROMPT -----")
            print("그 결과로 재실행: judge %s %s <기존 인자> subverdict=PASS|FAIL evidence=<핵심 인용요지>"
                  % (skill, cp))
            sys.exit(2)
        if sv == "FAIL":
            print("GATE FAIL [%s:%s]: substance 심판 FAIL — %s" % (skill, cp, args.get("evidence", "")))
            sys.exit(1)
        if not args.get("evidence", "").strip():
            print("GATE FAIL [%s:%s]: subverdict=PASS인데 evidence(인용) 비어있음 — 무효(근거 없는 PASS 거부)."
                  % (skill, cp))
            sys.exit(1)

    # (5) 통과 — verdict 토큰 기록(순서 강제용) + 이번 회차 form 출제 상태 소비(1회성)
    write_token(skill, cp, wf)
    try:
        os.remove(form_state_path(skill, cp, wf))
    except OSError:
        pass
    extra = " (substance VERIFIED)" if cpr.get("escalate") else ""
    print("GATE PASS [%s:%s]%s" % (skill, cp, extra))


def cmd_status(skill, kv_args):
    wf = _wf_key(skill, parse_kv(kv_args))
    tok = read_token(skill, wf)
    if tok is None:
        print("GATE STATUS [%s] wf=%s: (토큰 없음)" % (skill, wf))
    elif tok.get("_invalid"):
        print("GATE STATUS [%s] wf=%s: **무효 토큰** cp=%s ts=%s — %s"
              % (skill, wf, tok.get("cp"), tok.get("ts"), tok["_invalid"]))
    else:
        print("GATE STATUS [%s] wf=%s: 최근 통과 cp=%s ts=%s" % (skill, wf, tok.get("cp"), tok.get("ts")))


def cmd_clear(skill, kv_args):
    wf = _wf_key(skill, parse_kv(kv_args))
    _log_wipe("clear", skill, wf)
    remove_ledger(skill, wf)
    _clear_workflow(skill, wf)
    print("GATE CLEAR [%s] wf=%s — 이 워크플로의 토큰·감사큐 제거(워크플로 종료)." % (skill, wf))


# ---- audit (사후 룰베이스 감사 — 모델 자의 해석 0) ------------------------

def _walk(obj, acc):
    """JSONL 한 줄(obj)에서 tool_use(name,input)·tool_result(text)만 수집."""
    if isinstance(obj, dict):
        t = obj.get("type")
        if t == "tool_use":
            acc["uses"].append((obj.get("name", ""), obj.get("input", {}) or {}))
        elif t == "tool_result":
            c = obj.get("content")
            if isinstance(c, str):
                acc["results"].append(c)
            elif isinstance(c, list):
                for x in c:
                    if isinstance(x, dict) and "text" in x:
                        acc["results"].append(str(x.get("text", "")))
        for v in obj.values():
            _walk(v, acc)
    elif isinstance(obj, list):
        for v in obj:
            _walk(v, acc)


def _instr(inp):
    if not isinstance(inp, dict):
        return str(inp)
    out = []
    for v in inp.values():
        out.append(v if isinstance(v, str) else json.dumps(v, ensure_ascii=False))
    return " ".join(out)


def _chain_order(rubric):
    """루브릭의 'after' 링크로 cp 순서 복원(체인이 있을 때만). 단일 cp면 []."""
    afters = {cp: c.get("after") for cp, c in rubric.items()}
    if not any(afters.values()):
        return []
    order = []
    remaining = set(rubric.keys())
    progressed = True
    while remaining and progressed:
        progressed = False
        for cp in list(remaining):
            a = afters.get(cp)
            if a is None or a in order:
                order.append(cp); remaining.discard(cp); progressed = True
    order.extend(sorted(remaining))  # 순환 등 잔여는 뒤에
    return order


def cmd_audit(args):
    if not args:
        print("AUDIT 사용: audit <session.jsonl> | audit --latest [basedir]")
        sys.exit(3)
    if args[0] == "--latest":
        base = args[1] if len(args) > 1 else os.path.join(os.path.expanduser("~"), ".claude", "projects")
        cands = glob.glob(os.path.join(base, "**", "*.jsonl"), recursive=True)
        if not cands:
            print("AUDIT ERROR: .jsonl 세션 파일을 못 찾음(%s) — 경로를 직접 지정하라." % base)
            sys.exit(3)
        target = max(cands, key=os.path.getmtime)
    else:
        target = args[0]
    if not os.path.isfile(target):
        print("AUDIT ERROR: 파일 없음 %s" % target)
        sys.exit(3)

    rx_pass = re.compile(r"GATE PASS \[([^:\]]+):([^\]]+)\]")
    rx_fail = re.compile(r"GATE FAIL \[([^:\]]+):([^\]]+)\]")
    rx_needs = re.compile(r"GATE NEEDS-JUDGE \[([^:\]]+):([^\]]+)\]")
    rx_enq = re.compile(r"GATE ENQUEUED \[([^:\]]+):([^\]]+)\]")
    rx_rpass = re.compile(r"RESOLVE PASS \[([^\]]+)\]")
    rx_rfail = re.compile(r"RESOLVE FAIL \[([^\]]+)\]")
    rx_rhold = re.compile(r"RESOLVE HOLD \[([^\]]+)\]")
    gated_set = set(os.path.splitext(os.path.basename(p))[0]
                    for p in glob.glob(os.path.join(RUBRIC_DIR, "*.json")))
    invoked = set()        # 세션에서 활동 검출된 *게이트 대상* 스킬
    gate_pass = []; gate_fail = []; needs = []
    rubric_reads = []; rubric_writes = 0
    subagents = 0; subverdicts = []; judge_calls = 0; deliverables = []
    enqueued = []; resolve_pass = []; resolve_fail = []; resolve_hold = []; audit_records = 0

    with open(target, encoding="utf-8", errors="ignore") as fh:
        for ln, line in enumerate(fh, 1):
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
            except ValueError:
                continue
            acc = {"uses": [], "results": []}
            _walk(obj, acc)
            for (name, inp) in acc["uses"]:
                istr = _instr(inp)
                if _refs_sealed_rubric(istr):
                    if name in ("Read", "Grep"):
                        rubric_reads.append((ln, name))      # 봉인 정답지 열람 = 위반(audit_rubrics는 비봉인이라 제외)
                    elif name in ("Write", "Edit"):
                        rubric_writes += 1                   # 작성 = 정상
                    elif name in ("Bash", "PowerShell", "Shell") and _reads_via_shell(istr):
                        # ★ 셸로도 읽을 수 있다. `Read`/`Grep`만 보면 `cat`·`head`·
                        #   `python -c open(...)`으로 정답지를 그대로 열람할 수 있었다.
                        rubric_reads.append((ln, name + "(shell)"))
                if name == "Skill":
                    for g in gated_set:
                        if g in istr:
                            invoked.add(g)
                if name in ("Bash", "PowerShell", "Shell") and "stepgate.py" in istr:
                    for g in gated_set:
                        if g in istr:
                            invoked.add(g)
                if name in ("Agent", "Task"):
                    subagents += 1
                if name in ("Bash", "PowerShell", "Shell") and "stepgate.py" in istr \
                        and "judge" in istr and "audit-judge" not in istr:
                    judge_calls += 1
                    if "subverdict=PASS" in istr:
                        subverdicts.append(ln)
                if name in ("Bash", "PowerShell", "Shell") and "audit-judge" in istr and "--record" in istr:
                    audit_records += 1
                if name in ("Write", "Edit"):
                    fp = str(inp.get("file_path", "")) if isinstance(inp, dict) else ""
                    if is_deliverable_path(fp):
                        deliverables.append((ln, fp))
            for txt in acc["results"]:
                for m in rx_pass.finditer(txt):
                    gate_pass.append((ln, m.group(1), m.group(2)))
                for m in rx_fail.finditer(txt):
                    gate_fail.append((ln, m.group(1), m.group(2)))
                for m in rx_needs.finditer(txt):
                    needs.append((ln, m.group(1), m.group(2)))
                for m in rx_enq.finditer(txt):
                    enqueued.append((ln, m.group(1), m.group(2)))
                for m in rx_rpass.finditer(txt):
                    resolve_pass.append((ln, m.group(1)))
                for m in rx_rfail.finditer(txt):
                    resolve_fail.append((ln, m.group(1)))
                for m in rx_rhold.finditer(txt):
                    resolve_hold.append((ln, m.group(1)))

    findings = []
    passed_skills = set(sk for _, sk, _ in gate_pass)
    # R1 게이트 미호출 (게이트셋 인지 — 비게이트 스킬 산출물 오탐 제외)
    suspect = sorted(invoked - passed_skills)
    if suspect and deliverables:
        findings.append("[R1 게이트 미호출] 게이트 대상 스킬 %s 가 호출됐는데 GATE PASS 0 — judge를 안 부른 정황(산출물 %d건)." % (suspect, len(deliverables)))
    # R8 게이트 상태 초기화로 FAIL 탈출 — `init`/`clear`는 인증이 없으므로 **기록으로** 잡는다.
    wipes = []
    try:
        wp = os.path.join(gate_dir(), "wipe_log.jsonl")
        if os.path.isfile(wp):
            with open(wp, encoding="utf-8", errors="ignore") as fh:
                for line in fh:
                    try:
                        o = json.loads(line)
                    except ValueError:
                        continue
                    if o.get("token_cp") or o.get("pending_audits"):
                        wipes.append(o)
    except OSError:
        pass
    if gate_fail and wipes:
        findings.append(
            "[R8 상태 초기화] GATE FAIL %d건이 있고 `init`/`clear`가 **살아 있는 토큰·감사를 "
            "지운 기록** %d건(예 %s) — FAIL을 지우고 처음부터 다시 한 정황. 지운 것이 "
            "정당하면 사유를 남겨라."
            % (len(gate_fail), len(wipes),
               {k: wipes[-1].get(k) for k in ("ts", "action", "skill", "token_cp",
                                              "pending_audits")}))
    # R2 정답지 열람(Read/Grep on rubric)
    if rubric_reads:
        findings.append("[R2 정답지 열람] 봉인 루브릭을 Read/Grep %d건(라인 %s) — 합격 기준 열람 = 게이트 무효(Write=작성은 정상이라 제외)."
                        % (len(rubric_reads), ",".join(str(l) for l, _ in rubric_reads[:12])))
    # R3 FAIL 무시 진행
    if gate_fail and deliverables:
        last_fail = max(l for l, _, _ in gate_fail)
        deliv_after = [fp for l, fp in deliverables if l > last_fail]
        pass_after = [l for l, _, _ in gate_pass if l > last_fail]
        if deliv_after and not pass_after:
            findings.append("[R3 FAIL 무시] GATE FAIL(라인 %d) 이후 후속 PASS 없이 산출물 — 실패를 무시하고 진행한 정황." % last_fail)
    # R4 subverdict 위조 (substance 심판 없이 통과)
    if subverdicts and subagents == 0:
        findings.append("[R4 subverdict 위조] subverdict=PASS %d건인데 서브에이전트(Agent) 호출 0건 — 실물 심판 없이 통과시킨 정황." % len(subverdicts))
    # R5 순서 위반 (루브릭 after 체인에서 일반화 — 하드코딩 제거)
    for g in sorted(passed_skills):
        rb = load_rubric(g)
        if not rb:
            continue
        chain = _chain_order(rb)
        if not chain:
            continue
        passed_cps = set(cp for _, sk, cp in gate_pass if sk == g)
        idxs = [i for i, cp in enumerate(chain) if cp in passed_cps]
        if not idxs:
            continue
        last = max(idxs)
        missing = [chain[i] for i in range(last) if chain[i] not in passed_cps]
        if missing:
            findings.append("[R5 순서 위반] %s: '%s' PASS 이전 단계 누락 %s (체인 %s)." % (g, chain[last], missing, chain))
    # R6 비동기 감사 미완료(PENDING) — 디스크 큐를 세션 산출물과 대조(배경 에이전트 사망/미발사/미드레인)
    sess_delivs = set(os.path.normpath(fp).lower() for _, fp in deliverables)
    pend_hits = []
    for qf in glob.glob(os.path.join(audit_queue_dir(), "*.jsonl")):
        for _rid, o in read_queue_collapsed(qf).items():
            if o.get("status") == "PENDING":
                dp = os.path.normpath(o.get("deliverable", "")).lower()
                if dp and dp in sess_delivs:
                    pend_hits.append((o.get("skill"), o.get("cp"), o.get("deliverable")))
    if pend_hits:
        findings.append("[R6 비동기감사 미완료] 세션 산출물에 대한 PENDING 감사 %d건(예 %s) — 발사/기록 안 됨 또는 RESOLVE 미드레인."
                        % (len(pend_hits), pend_hits[0]))
    # R7 해소 게이트 누락 — 비동기 게이트(enqueue)·산출물은 있는데 RESOLVE PASS 없이 사인오프
    if enqueued and deliverables and not resolve_pass:
        findings.append("[R7 해소게이트 누락] GATE ENQUEUED %d건·산출물 %d건인데 RESOLVE PASS 0 — 감사를 resolve로 닫지 않고 사인오프한 정황."
                        % (len(enqueued), len(deliverables)))

    print("== 단계 게이트 사후 감사(룰베이스) [%s] ==" % os.path.basename(target))
    print("신호: PASS=%d FAIL=%d NEEDS-JUDGE=%d / judge호출=%d 서브에이전트=%d subverdict=PASS=%d / 루브릭 Read=%d(Write=%d) / 산출물=%d"
          % (len(gate_pass), len(gate_fail), len(needs), judge_calls, subagents, len(subverdicts), len(rubric_reads), rubric_writes, len(deliverables)))
    print("비동기: ENQUEUED=%d audit기록=%d / RESOLVE PASS=%d FAIL=%d HOLD=%d"
          % (len(enqueued), audit_records, len(resolve_pass), len(resolve_fail), len(resolve_hold)))
    print("게이트 대상 스킬 활동(호출 검출): %s" % (sorted(invoked) if invoked else "(없음 — 비게이트 세션이면 R1/R5 비적용)"))
    if gate_pass:
        print("통과 cp: " + ", ".join("%s:%s" % (s, c) for _, s, c in gate_pass))
    if not findings:
        print("AUDIT PASS — 룰 위반 0건. (한계: 통과한 무지성 도장·기록 안 된 행위는 사후로 탐지 못함.)")
        sys.exit(0)
    print("AUDIT FAIL — 룰 위반 %d건:" % len(findings))
    for f in findings:
        print("  - " + f)
    print(">>> 각 항목은 *결정론 규칙*의 탐지다. 해당 라인을 열어 확인하라(모델 추측 아님).")
    sys.exit(1)


# ---- 비동기 감사 커맨드 (T2/T3) ------------------------------------------

def cmd_enqueue(skill, cp, kv_args):
    args = parse_kv(kv_args)
    deliverable = args.get("deliverable", "").strip()
    if not deliverable:
        print("GATE ERROR: enqueue <skill> <cp> deliverable=<abspath> [dir=…] — deliverable 필수")
        sys.exit(3)
    reqid = _enqueue(skill, cp, args)
    _print_fire_block(skill, cp, reqid, deliverable)


def cmd_audit_judge(a):
    """audit-judge <reqid>            → 감사 패킷(lens·directives) 출력
       audit-judge <reqid> --record subverdict=PASS|FAIL findings=… evidence=…  → 결과 기록(FAIL이면 비0)."""
    if not a:
        print("AUDIT-JUDGE 사용: audit-judge <reqid> | audit-judge <reqid> --record subverdict=PASS|FAIL findings=… evidence=…")
        sys.exit(3)
    reqid = a[0]
    rest = a[1:]
    path, entry = _find_queue_entry(reqid)
    if entry is None:
        # 복구: 배경 감사 도중 부모가 같은 워크플로를 clear/init/재빌드(supersede)해 큐 파일이
        # 사라졌어도 감사 결과를 잃지 않도록, 신형 reqid에서 skill/cp/wf를 복원해 엔트리를 재생성한다.
        # (deliverable/deliv_hash는 모름 → 빈 값. resolve의 deliverable-scoping은 정상 엔트리를
        #  우선하므로 이 복구 엔트리가 스코프된 해소를 오염시키지 않는다.)
        rec = _reqid_parts(reqid)
        if rec is None:
            print("AUDIT-JUDGE ERROR: reqid '%s' 큐에 없음 + 복구 불가(구형 형식) — enqueue 먼저." % reqid)
            sys.exit(3)
        rskill, rcp, rwf = rec
        path = queue_path(rskill, rwf)
        entry = {"reqid": reqid, "skill": rskill, "cp": rcp, "deliverable": "",
                 "deliv_hash": "", "wfid": rwf, "dir": "", "status": "PENDING",
                 "subverdict": None, "findings": "", "evidence": "",
                 "ts_enq": None, "ts_done": None, "agent": None, "recovered": True}
        print("AUDIT-JUDGE [복구] reqid '%s' 큐 파일 부재(clear/재빌드로 유실 추정) — reqid에서 엔트리 재생성해 기록 계속." % reqid)
    skill = entry.get("skill"); cp = entry.get("cp"); deliverable = entry.get("deliverable", "")
    if "--record" in rest:
        kv = parse_kv([x for x in rest if x != "--record"])
        sv = kv.get("subverdict", "").strip().upper()
        if sv not in ("PASS", "FAIL"):
            print("AUDIT-JUDGE ERROR: --record엔 subverdict=PASS|FAIL 필요.")
            sys.exit(3)
        if sv == "FAIL" and not (kv.get("findings", "").strip() or kv.get("evidence", "").strip()):
            print("AUDIT-JUDGE ERROR: subverdict=FAIL엔 findings/evidence(위반 인용) 필요 — 근거 없는 FAIL 무효.")
            sys.exit(3)
        # ★ PASS도 근거를 요구한다. 예전에는 FAIL만 인용을 요구했다 — 비대칭이 거꾸로였다.
        #   이 시스템에서 위험한 것은 엄한 FAIL이 아니라 **근거 없는 관대한 PASS**이고,
        #   PASS가 공짜면 감사자는 산출물을 열지 않아도 통과를 기록할 수 있다.
        #   그래서 `evidence`가 **산출물 안에 실제로 있는 문구**여야 한다(열어봤다는 증거).
        if sv == "PASS":
            ev = kv.get("evidence", "").strip()
            if not ev:
                print("AUDIT-JUDGE ERROR: subverdict=PASS에도 evidence(산출물에서 그대로 뽑은 "
                      "문구) 필요 — 열어보지 않은 PASS는 인정하지 않는다.")
                sys.exit(3)
            if deliverable and os.path.isfile(deliverable):
                ok, why = quote_check(deliverable, ev, 20)
                if not ok:
                    print("AUDIT-JUDGE ERROR: evidence가 산출물에 없다 — %s\n"
                          "  PASS를 기록하려면 그 파일에서 **그대로** 뽑은 20자 이상 문구를 내라."
                          % why)
                    sys.exit(3)
        # ★ 기록된 FAIL은 덮어쓸 수 없다. 큐가 append-only이고 읽을 때 마지막 줄만 접으므로,
        #   예전에는 FAIL 뒤에 PASS를 한 줄 더 쓰면 FAIL이 사라졌다. 수정했으면
        #   **새 vN_M으로 새 reqid를 받아 재감사**하는 것이 규약이다(같은 reqid 갱신이 아니다).
        cur = read_queue_collapsed(path).get(reqid) or {}
        if str(cur.get("status")) == "FAIL" and sv == "PASS":
            print("AUDIT-JUDGE ERROR: 이 reqid는 이미 FAIL로 기록됐다 — PASS로 덮을 수 없다.\n"
                  "  산출물을 새 vN_M으로 고치고 그 cp를 다시 judge해 **새 감사**를 받아라.\n"
                  "  기존 위반: %s" % (str(cur.get("findings") or cur.get("evidence"))[:200]))
            sys.exit(3)
        upd = dict(entry)
        upd["status"] = sv
        upd["subverdict"] = sv
        upd["findings"] = kv.get("findings", "")
        upd["evidence"] = kv.get("evidence", "")
        upd["ts_done"] = time.strftime("%Y-%m-%d %H:%M:%S")
        upd["agent"] = "recorded"
        append_queue(path, upd)
        print("AUDIT RECORD [%s:%s] req=%s → %s" % (skill, cp, reqid, sv))
        if sv == "FAIL":
            print("  위반: %s" % kv.get("findings", "") or kv.get("evidence", ""))
            sys.exit(1)
        sys.exit(0)
    # packet 모드
    ar = load_audit_rubric(skill)
    print("== AUDIT PACKET [%s:%s] req=%s ==" % (skill, cp, reqid))
    print("산출물: %s" % deliverable)
    if entry.get("dir"):
        print("작업 폴더: %s" % entry["dir"])
    if not ar or cp not in ar:
        print("[경고] audit_rubric 없음(_stepgate/audit_rubrics/%s.json[%s]) — 기본 빡빡 검수: "
              "스킬 가이드 전반을 산출물과 *그대로 인용* 대조, 인용 못 하면 PASS." % (skill, cp))
    elif not (ar[cp].get("lenses") or []):
        # ★ cp는 있는데 렌즈가 0개면 감사자가 볼 것이 없다 — 그러면 **아무 검사도 없이
        #   PASS**가 나온다. "렌즈가 없다"는 통과가 아니라 **감사 정의 오류**다.
        print("AUDIT ERROR: audit_rubric[%s][%s]에 lens가 0개다 — 검사할 것이 없는 감사는 "
              "통과가 아니라 정의 오류다. lens를 채우거나 cp를 지워 기본 검수로 넘겨라."
              % (skill, cp))
        sys.exit(3)
    else:
        cpr = ar[cp]
        print("아래 lens를 *각각* 판정하라. 규칙: 위반 라인/셀좌표/캡션을 *그대로 인용* 못 하면 그 lens는 PASS(추측 FAIL 무효).")
        for d in cpr.get("global_directives", []):
            print("  ※ %s" % d)
        for L in cpr.get("lenses", []):
            nf = " [needs_fetch: 링크 실제 WebFetch]" if L.get("needs_fetch") else ""
            print("  - lens[%s] (%s)%s: %s" % (L.get("id", "?"), L.get("severity", "?"), nf, L.get("ask", "")))
    print("종료: SG audit-judge %s --record subverdict=PASS|FAIL findings=<위반 요약(없으면 빈)> evidence=<핵심 인용>" % reqid)


def _onedrive_sync(folder):
    """RESOLVE PASS(사인오프) 직후 그 산출물 폴더를 OneDrive(7분류 재구성본)로 자동 반영.
    best-effort — 싱크 실패/미매핑이어도 사인오프는 막지 않는다(예외 삼킴). dev_root 밖/미매핑은
    sync_onedrive.py가 알아서 스킵. 산출물+마스터md만, 정크(_raw_sources·이전버전·메모) 제외."""
    here = os.path.dirname(os.path.abspath(__file__))
    script = os.path.join(here, "sync_onedrive.py")
    if not folder or not os.path.isfile(script):
        return
    try:
        import subprocess
        r = subprocess.run([sys.executable, script, folder],
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace", timeout=120)
        print("── OneDrive 싱크(사인오프 자동) ──")
        out = (r.stdout or "").strip()
        print(out if out else "(출력 없음)")
        if r.returncode == 2:
            print("  [주의] 매핑 미등록 폴더 — OneDrive 배치 미정(knowledge-map start로 카테고리 등록 후 재싱크).")
    except Exception as e:
        print("── OneDrive 싱크 건너뜀(best-effort 실패, 사인오프엔 영향 없음): %s ──" % e)


def cmd_resolve(skill, kv_args):
    """T3 해소 게이트 — 사인오프 직전. 이 워크플로의 활성 감사가 전부 PASS여야 RESOLVE PASS.
    PASS 시 그 산출물 폴더를 OneDrive로 자동 싱크(_onedrive_sync, best-effort)."""
    args = parse_kv(kv_args)
    deliverable = args.get("deliverable", "").strip()
    wf = _wf_key(skill, args)
    coll = read_queue_collapsed(queue_path(skill, wf))
    active = [o for o in coll.values() if o.get("status") != "SUPERSEDED"]
    # deliverable 지정 시 *그 산출물의* 감사만 해소 대상 — 같은 폴더 타 산출물 PENDING이 막지 않게.
    if deliverable:
        dh = deliv_hash(deliverable)
        scoped = [o for o in active if o.get("deliv_hash") == dh]
        if scoped:
            active = scoped
    if not active:
        print("RESOLVE FAIL [%s]: 이 워크플로/산출물에 적재된 비동기 감사 0건 — 게이트 미경유 산출물"
              "(judge async로 enqueue 안 됨)." % skill)
        sys.exit(1)
    pend = [o for o in active if o.get("status") == "PENDING"]
    fail = [o for o in active if o.get("status") == "FAIL"]
    if pend:
        print("RESOLVE HOLD [%s]: 비동기 감사 %d건 미완료(PENDING) — 사인오프 전 지금 동기 드레인하라:" % (skill, len(pend)))
        for o in pend:
            print("  · req=%s cp=%s 산출물=%s → SG audit-judge %s (검수 후 --record)"
                  % (o.get("reqid"), o.get("cp"), o.get("deliverable"), o.get("reqid")))
        sys.exit(2)
    if fail:
        print("RESOLVE FAIL [%s]: 비동기 감사 %d건 FAIL — 새 vN_M로 수정·재빌드 후 해당 cp judge async 재실행(재감사):" % (skill, len(fail)))
        for o in fail:
            print("  · req=%s cp=%s: %s" % (o.get("reqid"), o.get("cp"), o.get("findings") or o.get("evidence")))
        sys.exit(1)
    cps = sorted(set(str(o.get("cp")) for o in active))
    print("RESOLVE PASS [%s] — 활성 감사 %d건 전부 PASS (cp: %s). 이제 사용자에게 done 보고 가능."
          % (skill, len(active), ", ".join(cps)))
    if deliverable and not any(o.get("deliverable") == deliverable for o in active):
        print("  [주의] 지정 deliverable=%s 에 대한 감사 엔트리 없음(경로 상이) — 확인 요망." % deliverable)
    # 사인오프 자동화 — 보고 준비된(전 감사 PASS) 산출물 폴더를 OneDrive로 반영(best-effort).
    sync_target = os.path.dirname(deliverable) if deliverable else args.get("dir", "")
    _onedrive_sync(sync_target)
    sys.exit(0)


def cmd_audit_status(skill, kv_args):
    wf = _wf_key(skill, parse_kv(kv_args))
    coll = read_queue_collapsed(queue_path(skill, wf))
    active = [o for o in coll.values() if o.get("status") != "SUPERSEDED"]
    tally = {"PENDING": 0, "PASS": 0, "FAIL": 0}
    for o in active:
        st = o.get("status", "PENDING")
        tally[st] = tally.get(st, 0) + 1
    print("AUDIT STATUS [%s] wfid=%s: PENDING=%d PASS=%d FAIL=%d (활성 %d)"
          % (skill, wf, tally.get("PENDING", 0), tally.get("PASS", 0),
             tally.get("FAIL", 0), len(active)))
    for o in active:
        print("  · %s cp=%s 산출물=%s" % (o.get("status"), o.get("cp"), o.get("deliverable")))


def main():
    a = sys.argv[1:]
    if len(a) < 2:
        print(__doc__)
        sys.exit(3)
    cmd, skill = a[0], a[1]
    if cmd == "init":
        cmd_init(skill, a[2:])
    elif cmd == "form":
        if len(a) < 3:
            print("GATE ERROR: form <skill> <cp> [dir=<과제폴더>|deliverable=<abspath>]")
            sys.exit(3)
        cmd_form(skill, a[2], a[3:])
    elif cmd == "judge":
        if len(a) < 3:
            print("GATE ERROR: judge <skill> <cp> [<id>=<값> ...]")
            sys.exit(3)
        cmd_judge(skill, a[2], a[3:])
    elif cmd == "status":
        cmd_status(skill, a[2:])
    elif cmd == "clear":
        cmd_clear(skill, a[2:])
    elif cmd == "audit":
        cmd_audit(a[1:])
    elif cmd == "enqueue":
        if len(a) < 3:
            print("GATE ERROR: enqueue <skill> <cp> deliverable=<abspath> [dir=…]")
            sys.exit(3)
        cmd_enqueue(skill, a[2], a[3:])
    elif cmd == "audit-judge":
        cmd_audit_judge(a[1:])
    elif cmd == "resolve":
        cmd_resolve(skill, a[2:])
    elif cmd == "audit-status":
        cmd_audit_status(skill, a[2:])
    else:
        print("GATE ERROR: unknown cmd '%s'" % cmd)
        sys.exit(3)


if __name__ == "__main__":
    main()
