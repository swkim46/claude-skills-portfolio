# -*- coding: utf-8 -*-
"""research-docx 스킬 — python-docx 하우스스타일 헬퍼.
references/docx_consulting_report_guide.md(스킬 동봉) 준거. 맑은 고딕·하이퍼링크 인용·표 출처 링크·Q/B/P/TABLE.

사용 패턴:
    import docx_helpers as D
    doc = D.new_doc()
    SRC = {'R1':'https://...', 'C2':'https://...'}
    D.H1(doc, '1. 현황')
    D.B(doc, [('인구(2025): ', True), ('약 284M명', False)], refs=D.R(SRC,'R1'))
    D.TABLE(doc, ['지표','값','출처'], [['NIM','4.54%','N4']], src_col=2, SRC=SRC)
    D.Q(doc, 'verbatim 원문', 'R1', SRC['R1'])
    doc.save('out.docx')
"""
import re
import os
import shutil
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = '맑은 고딕'
GREY = RGBColor(0x70, 0x70, 0x70)
DARK = RGBColor(0x1F, 0x1F, 0x1F)
LINKBLUE = '0563C1'
_LABEL_RE = re.compile(r'([A-Za-z]+\d+[a-z]?)')


def new_doc(base_font=FONT, base_size=10, margins_cm=(2.0, 2.0, 2.2, 2.2)):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = base_font
    st._element.rPr.rFonts.set(qn('w:eastAsia'), base_font)
    st.font.size = Pt(base_size)
    for s in doc.sections:
        s.top_margin = Cm(margins_cm[0]); s.bottom_margin = Cm(margins_cm[1])
        s.left_margin = Cm(margins_cm[2]); s.right_margin = Cm(margins_cm[3])
    return doc


def _font(r, name=FONT, size=10, bold=False, italic=False, color=None):
    r.font.name = name
    r._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    if color:
        r.font.color.rgb = color


def add_hyperlink(paragraph, url, text, font_size=9, font_name=FONT,
                  is_eastasia=True, color_hex=LINKBLUE, bold=False):
    """본문/표에 클릭 가능한 외부 하이퍼링크 삽입(word/_rels에 관계 등록)."""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
    run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rF = OxmlElement('w:rFonts'); rF.set(qn('w:ascii'), font_name); rF.set(qn('w:hAnsi'), font_name)
    if is_eastasia:
        rF.set(qn('w:eastAsia'), font_name)
    rPr.append(rF)
    c = OxmlElement('w:color'); c.set(qn('w:val'), color_hex); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(font_size * 2))); rPr.append(sz)
    if bold:
        rPr.append(OxmlElement('w:b'))
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve')
    run.append(t); h.append(run); paragraph._p.append(h)
    return h


def title(doc, text, sub=None, note=None):
    p = doc.add_paragraph(); _font(p.add_run(text), size=18, bold=True)
    if sub:
        s = doc.add_paragraph(); _font(s.add_run(sub), size=9, color=GREY)
    if note:
        n = doc.add_paragraph(); n.paragraph_format.space_after = Pt(6)
        _font(n.add_run(note), size=8.5, color=GREY, italic=True)


def H1(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(8)
    _font(p.add_run(text), size=14, bold=True, color=DARK)

def H2(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    _font(p.add_run(text), size=11.5, bold=True)

def H3(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    _font(p.add_run(text), size=10.5, bold=True)


def _refs(p, refs, size):
    for label, url in refs:
        _font(p.add_run(' ['), size=size - 1, color=RGBColor(0x60, 0x60, 0x60))
        add_hyperlink(p, url, label, font_size=size - 1)
        _font(p.add_run(']'), size=size - 1, color=RGBColor(0x60, 0x60, 0x60))


def P(doc, parts, size=10, space_after=4):
    """parts: str | [(text, bold), ...]"""
    if isinstance(parts, str):
        parts = [(parts, False)]
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.line_spacing = 1.4
    for text, bold in parts:
        _font(p.add_run(text), size=size, bold=bold)
    return p


def B(doc, parts, level=0, size=10, refs=None):
    """불릿. refs=[(label,url)] → 끝에 클릭 가능한 [label]."""
    if isinstance(parts, str):
        parts = [(parts, False)]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.7); p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.35
    bullet = '•' if level == 0 else ('◦' if level == 1 else '–')
    _font(p.add_run(f'{bullet}  '), size=size)
    for text, bold in parts:
        _font(p.add_run(text), size=size, bold=bold)
    if refs:
        _refs(p, refs, size)
    return p


def Q(doc, quote, source_label=None, source_url=None):
    """원문 verbatim 인용(이탤릭 들여쓰기 + 라벨 링크)."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.35
    _font(p.add_run(f'"{quote}"'), size=9.5, italic=True, color=RGBColor(0x33, 0x33, 0x33))
    if source_label and source_url:
        _font(p.add_run('  ['), size=8.5, color=RGBColor(0x60, 0x60, 0x60))
        add_hyperlink(p, source_url, source_label, font_size=8.5)
        _font(p.add_run(']'), size=8.5, color=RGBColor(0x60, 0x60, 0x60))
    return p


def _fill_src_cell(cell, spec, SRC, size=8.5):
    """표 출처 셀: 'C1·R2 / O4' 등 라벨을 각각 클릭 하이퍼링크로(구분자는 평문)."""
    cell.text = ''
    para = cell.paragraphs[0]
    pos = 0
    for m in _LABEL_RE.finditer(spec):
        if m.start() > pos:
            _font(para.add_run(spec[pos:m.start()]), size=size)
        label = m.group(1)
        if SRC and label in SRC:
            add_hyperlink(para, SRC[label], label, font_size=size)
        else:
            _font(para.add_run(label), size=size)
        pos = m.end()
    if pos < len(spec):
        _font(para.add_run(spec[pos:]), size=size)


def _emit_md_runs(para, text, size, bold=False):
    """셀/문단 텍스트의 **굵게** 구간을 bold run으로 렌더(나머지 평문). 균형 잡힌 ** 가정."""
    for idx, seg in enumerate(re.split(r'\*\*', text)):
        if seg == '':
            continue
        _font(para.add_run(seg), size=size, bold=(bold or idx % 2 == 1))


def TABLE(doc, headers, rows, src_col=None, SRC=None, style='Light Grid Accent 1', size=9, md_bold=False, col_widths=None):
    """src_col: 0-based 열 인덱스 → 그 열 셀의 출처 라벨을 클릭 하이퍼링크로(SRC 필요).
    md_bold=True면 비-출처 셀의 **텍스트**를 굵게 렌더(opt-in; 기존 호출 영향 없음).
    col_widths: [cm, ...] 0-based 열별 고정 폭(autofit off + tblLayout fixed + 셀 width).
        지정 시 Word 자동분배를 끄고 폭을 고정 — 출처 등 짧은 열이 과도하게 넓어지는 것 방지.
        합은 본문 폭(A4 좌우여백 2.2cm 기준 약 16.6cm) 이하로. None이면 기존 동작(opt-in)."""
    t = doc.add_table(rows=1 + len(rows), cols=len(headers)); t.style = style
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]; cell.text = ''
        _font(cell.paragraphs[0].add_run(h), size=size, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[1 + ri].cells[ci]
            if src_col is not None and ci == src_col:
                _fill_src_cell(cell, str(val), SRC, size=size - 0.5)
            elif md_bold and '**' in str(val):
                cell.text = ''
                _emit_md_runs(cell.paragraphs[0], str(val), size)
            else:
                cell.text = ''
                _font(cell.paragraphs[0].add_run(str(val)), size=size)
    if col_widths:
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = tblPr.find(qn('w:tblLayout'))
        if lay is None:
            lay = OxmlElement('w:tblLayout'); tblPr.append(lay)
        lay.set(qn('w:type'), 'fixed')
        for r in t.rows:
            for ci, w in enumerate(col_widths):
                if ci < len(r.cells):
                    r.cells[ci].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def R(SRC, *labels):
    """refs 헬퍼: D.R(SRC,'R1','R2') → [('R1',url),('R2',url)]."""
    return [(l, SRC[l]) for l in labels]


def src_section(doc, title_text, items, SRC):
    """출처 섹션: items=[(label, desc, key)] → 라벨 하이퍼링크 + URL 줄(Consolas).
    Tier 등급은 클라이언트 산출물에서 제외(헤더 텍스트에 (Tier ..) 쓰지 말 것)."""
    H2(doc, title_text)
    for label, desc, key in items:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.3
        _font(p.add_run('•  '), size=9)
        add_hyperlink(p, SRC[key], f'[{label}] {desc}', font_size=9)
        p2 = doc.add_paragraph(); p2.paragraph_format.left_indent = Cm(0.8); p2.paragraph_format.space_after = Pt(3); p2.paragraph_format.line_spacing = 1.2
        add_hyperlink(p2, SRC[key], SRC[key], font_size=7.5, font_name='Consolas', is_eastasia=False, color_hex='404080')


def verify(path):
    """빌드 후 검증: 외부 하이퍼링크 수 / unique URL."""
    import zipfile
    with zipfile.ZipFile(path) as z:
        rels = z.read('word/_rels/document.xml.rels').decode()
    ext = len(re.findall('TargetMode="External"', rels))
    urls = len(set(re.findall(r'Target="(https?://[^"]+)"', rels)))
    return {'external_links': ext, 'unique_urls': urls}


# ── 자동 버저닝 (재빌드 덮어쓰기 근본 차단) ──────────────────────────────
# 고정 파일명에 저장하면 재빌드마다 덮어써진다. 아래 헬퍼는 폴더를 스캔해
# 다음 vN_M을 자동 계산하고, 그 경로가 이미 있으면 예외를 던진다(clobber-proof).
# major/minor 판단만 호출부(사람)가 하고, 번호·이관은 도구가 자동 처리한다.
# cf. [[deliverable-versioning]] 이관 2규칙.
_VERPATH_RE = re.compile(r'_v(\d+)_(\d+)\.[A-Za-z0-9]+$')


def next_version_path(folder, stem, major=False, ext='docx'):
    """다음 `{stem}_vN_M.{ext}` 경로를 계산해 반환(순수 — 파일 이동 없음).
    folder를 스캔해:
      - 버전 없음 → v1_0
      - major=False(기본): 현재 최고 major N의 *최신 minor + 1* → v{N}_{M+1}
      - major=True: v{N+1}_0  (구조·결론·핵심수치 변경 시 — 호출부가 판단)
    반환 경로가 이미 존재하면 FileExistsError(덮어쓰기 방지). 저장은 호출부(doc.save)
    또는 save_versioned()가 한다. (major/minor는 도구가 자동판별 불가한 의미판단이라
    인자로 받는다 — 기본 minor라 미지정 시 안전한 쪽(구버전 보존)으로 떨어짐.)"""
    vers = []
    prefix = stem + '_v'
    if os.path.isdir(folder):
        for fn in os.listdir(folder):
            if fn.startswith(prefix) and fn.endswith('.' + ext):
                m = _VERPATH_RE.search(fn)
                if m:
                    vers.append((int(m.group(1)), int(m.group(2))))
    if not vers:
        n, mm = 1, 0
    elif major:
        n, mm = max(v[0] for v in vers) + 1, 0
    else:
        n = max(v[0] for v in vers)
        mm = max(m2 for n2, m2 in vers if n2 == n) + 1
    path = os.path.join(folder, "%s_v%d_%d.%s" % (stem, n, mm, ext))
    if os.path.exists(path):
        raise FileExistsError("clobber-proof: %s already exists" % path)
    return path


def archive_superseded(folder, stem, ext='docx'):
    """이관 2규칙: 각 major는 *최신 minor 1개만* 메인에 남기고, 더 낮은 minor는
    folder/이전버전/ 으로 이동(구버전 보존·삭제 아님). major 계열은 메인에 누적."""
    prefix = stem + '_v'
    by_n = {}
    if not os.path.isdir(folder):
        return
    for fn in os.listdir(folder):
        if fn.startswith(prefix) and fn.endswith('.' + ext):
            m = _VERPATH_RE.search(fn)
            if m:
                by_n.setdefault(int(m.group(1)), []).append((int(m.group(2)), fn))
    prevdir = os.path.join(folder, '이전버전')
    for n, lst in by_n.items():
        if len(lst) < 2:
            continue
        latest = max(m2 for m2, _ in lst)
        for m2, fn in lst:
            if m2 == latest:
                continue
            os.makedirs(prevdir, exist_ok=True)
            src, dst = os.path.join(folder, fn), os.path.join(prevdir, fn)
            if os.path.exists(src) and not os.path.exists(dst):
                try:
                    shutil.move(src, dst)
                except (PermissionError, OSError) as e:
                    # 구버전이 잠김(WinError 32 등)이면 이관은 best-effort로 건너뜀 —
                    # 저장은 이미 끝났으므로 빌드를 크래시시키지 않는다(나중에 수동 정리).
                    print("  [archive skip] %s 이관 건너뜀(잠김 등): %s" % (fn, e))


def save_versioned(doc, folder, stem, major=False, ext='docx'):
    """자동 버저닝 저장: next_version_path()로 새 vN_M 경로 계산 → doc.save →
    직전 minor 이관(archive_superseded). 재빌드마다 항상 새 파일이라 in-place
    덮어쓰기가 구조적으로 불가능하다. 반환: 저장한 경로(str).
    사용: path = D.save_versioned(doc, r'<REPO>\\<과제>', '<stem>')  # 기본 minor
          path = D.save_versioned(doc, folder, stem, major=True)             # 주요 변경"""
    os.makedirs(folder, exist_ok=True)
    path = next_version_path(folder, stem, major=major, ext=ext)
    doc.save(path)
    # 이관은 best-effort — 어떤 이유로든 실패해도 저장된 새 버전 경로 반환을 막지 않는다.
    try:
        archive_superseded(folder, stem, ext=ext)
    except Exception as e:
        print("  [archive skip] 이관 단계 오류(저장은 성공): %s" % e)
    return path
