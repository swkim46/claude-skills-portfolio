#!/usr/bin/env python3
"""Build (or append to) the expert-VQ compilation workbook from a JSON spec.

Layout (transposed response block; top questions table NOT shifted):
  Col A    -> narrow empty margin (width ~2.5)
  Col B    -> VQ No. (top questions table only; width ~8)
  Col C    -> Questions text (top table) AND field names (response block); width ~22
  Col D+   -> one wide column per expert (width ~90)

Freeze panes are computed (never hardcoded) from the detected last question row.

MULTI-SHEET behaviour (questions changed -> new sheet):
  When appending to an existing workbook, the script compares the spec's
  questions against each existing response sheet:
    * If a sheet's questions MATCH -> the new experts are appended to that sheet.
    * If NO sheet matches (the VQ set is different) -> a NEW response sheet is
      created (e.g. "VQ 답변 모음 (2)") with its own questions table, PRECEDED by a
      topic divider tab: an empty, name-only sheet named from project_label
      (e.g. "0616(과제 A) >>") — same role as the first round's divider.
  A brand-new file gets the divider tab + the first response sheet.
  => Always pass project_label (the round's topic, parsed from the email) so the
     new sheet gets its divider.

Usage:
    python build_vq_excel.py spec.json output.xlsx

Spec shape:
{
  "project_label": "0615(해외 증권사 C, 해외 은행 D) >>",
  "sheet_name": "VQ 답변 모음",        # optional; base name for the response sheet
  "questions": [{"short": "Q1. ...", "full": "...full question text..."}, ...],
  "experts": [
    {"network": "thirdbridge", "name": "#11.2 - Fabien Legland",
     "background": "Current Executive Director ... (10/2022 - Present)",
     "answers": ["...Q1 (orig + blank line + KR)...", ...],
     "availability": "평일 오전 9~11시 (보통 가능)", "restricted": ""}
  ]
}
"""
import json
import os
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

RESP_SHEET = "VQ 답변 모음"
FONT = "Calibri"

# Column geometry (1-indexed)
MARGIN_COL = 1            # A
VQNO_COL = 2             # B
LABEL_COL = 3            # C
FIRST_DATA_COL = 4       # D

MARGIN_WIDTH = 2.5
VQNO_WIDTH = 8
LABEL_WIDTH = 22
DATA_WIDTH = 90

QUESTIONS_START_ROW = 3

FIXED_LABELS = {
    "network": "전문가 네트워크 이름",
    "name": "전문가 이름",
    "background": "주요 경력",
    "availability": "가능시간",
    "restricted": "제한 주제",
}


def _field_keys(n_q):
    return ["network", "name", "background"] + [f"q{i}" for i in range(n_q)] + \
           ["availability", "restricted"]


def _labels(questions):
    lab = dict(FIXED_LABELS)
    for i, q in enumerate(questions):
        lab[f"q{i}"] = q.get("short", f"Q{i+1}")
    return lab


def _expert_value(e, key):
    if key in ("network", "name", "background", "availability", "restricted"):
        return e.get(key, "")
    if key.startswith("q"):
        i = int(key[1:])
        ans = e.get("answers", [])
        return ans[i] if i < len(ans) else ""
    return ""


def _norm(s):
    return " ".join(str(s or "").split())


def _is_response_sheet(ws):
    return ws.cell(row=2, column=VQNO_COL).value == "VQ No."


def _read_sheet_questions(ws):
    """Full question texts from the top table (column C, rows 3..)."""
    qs = []
    r = QUESTIONS_START_ROW
    while ws.cell(row=r, column=VQNO_COL).value not in (None, ""):
        qs.append(ws.cell(row=r, column=LABEL_COL).value or "")
        r += 1
    return qs


def _questions_match(spec_questions, existing_texts):
    spec_texts = [_norm(q.get("full", q.get("short", ""))) for q in spec_questions]
    return spec_texts == [_norm(t) for t in existing_texts]


def _read_experts_from_sheet(ws):
    start = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=LABEL_COL).value == FIXED_LABELS["network"]:
            start = r
            break
    if start is None:
        return []
    label_rows = {}
    for r in range(start, ws.max_row + 1):
        lbl = ws.cell(row=r, column=LABEL_COL).value
        if lbl in (None, ""):
            break
        label_rows[lbl] = r
    q_rows = sorted(r for lbl, r in label_rows.items()
                    if lbl not in FIXED_LABELS.values())
    experts = []
    col = FIRST_DATA_COL
    while True:
        net = ws.cell(row=label_rows[FIXED_LABELS["network"]], column=col).value
        name = ws.cell(row=label_rows[FIXED_LABELS["name"]], column=col).value
        if (net in (None, "")) and (name in (None, "")):
            break
        experts.append({
            "network": net or "",
            "name": name or "",
            "background": ws.cell(row=label_rows[FIXED_LABELS["background"]], column=col).value or "",
            "answers": [ws.cell(row=r, column=col).value or "" for r in q_rows],
            "availability": ws.cell(row=label_rows[FIXED_LABELS["availability"]], column=col).value or "",
            "restricted": ws.cell(row=label_rows[FIXED_LABELS["restricted"]], column=col).value or "",
        })
        col += 1
    return experts


def _unique_sheet_name(wb, base):
    base = base[:31]
    if base not in wb.sheetnames:
        return base
    i = 2
    while True:
        cand = f"{base} ({i})"[:31]
        if cand not in wb.sheetnames:
            return cand
        i += 1


def _force_text_leading_formula(ws):
    """VQ 트래커는 전부 텍스트 셀 — 답변이 '=', '+', '-', '@'로 시작하면(예: 이메일의 '==>' 답변 마커)
    openpyxl이 그 셀을 formula로 저장하고, Excel이 열 때 파싱 실패로 파일을 손상시킨다
    ('Removed Records: Formula' 복구 팝업). 이 시트엔 의도된 수식이 전혀 없으므로 그런 셀을
    모두 강제 텍스트(data_type='s')로 만든다. 실사례: #25.1 답변 4칸('==> ...')이 formula로
    저장돼 손상 → 이 가드로 재발 차단."""
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
                c.data_type = "s"


def _write_sheet(ws, questions, experts):
    """Write the full transposed sheet content into a blank worksheet."""
    n_q = len(questions)
    keys = _field_keys(n_q)
    labels = _labels(questions)

    bold = Font(name=FONT, size=11, bold=True)
    reg = Font(name=FONT, size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    nowrap = Alignment(wrap_text=False, vertical="top")

    # Top questions table: B = VQ No., C = Questions (wrap OFF)
    h1 = ws.cell(row=2, column=VQNO_COL, value="VQ No."); h1.font = bold; h1.alignment = nowrap
    h2 = ws.cell(row=2, column=LABEL_COL, value="Questions"); h2.font = bold; h2.alignment = nowrap
    last_q_row = QUESTIONS_START_ROW - 1
    for i, q in enumerate(questions):
        r = QUESTIONS_START_ROW + i
        last_q_row = r
        n = ws.cell(row=r, column=VQNO_COL, value=i + 1); n.font = reg; n.alignment = nowrap
        cc = ws.cell(row=r, column=LABEL_COL, value=q.get("full", q.get("short", "")))
        cc.font = reg; cc.alignment = nowrap

    # Transposed response block: labels in C, experts in D+
    block_start = last_q_row + 2          # exactly one blank row gap (computed)
    for ridx, key in enumerate(keys):
        r = block_start + ridx
        lc = ws.cell(row=r, column=LABEL_COL, value=labels[key])
        lc.font = bold; lc.alignment = wrap
        for cidx, e in enumerate(experts):
            dc = ws.cell(row=r, column=FIRST_DATA_COL + cidx, value=_expert_value(e, key))
            dc.font = reg; dc.alignment = wrap

    # Widths
    ws.column_dimensions[get_column_letter(MARGIN_COL)].width = MARGIN_WIDTH
    ws.column_dimensions[get_column_letter(VQNO_COL)].width = VQNO_WIDTH
    ws.column_dimensions[get_column_letter(LABEL_COL)].width = LABEL_WIDTH
    for cidx in range(max(len(experts), 1)):
        ws.column_dimensions[get_column_letter(FIRST_DATA_COL + cidx)].width = DATA_WIDTH

    # Freeze panes (computed, never hardcoded): body's first data cell
    ws.freeze_panes = f"{get_column_letter(FIRST_DATA_COL)}{block_start}"
    _force_text_leading_formula(ws)   # '=' 시작 답변이 formula로 저장돼 Excel 파일 손상되는 것 방지
    return block_start


def build(spec, out_path):
    questions = spec.get("questions", [])
    new_experts = spec.get("experts", [])
    base_name = (spec.get("sheet_name") or RESP_SHEET)
    divider_name = None                                  # 토픽 구분(divider) 탭 — 이름만 의미, 내용 없음

    if os.path.exists(out_path):
        wb = load_workbook(out_path)
        target_name = None
        for sn in wb.sheetnames:
            ws = wb[sn]
            if _is_response_sheet(ws) and _questions_match(questions, _read_sheet_questions(ws)):
                target_name = sn
                break
        if target_name is not None:                      # questions match -> append
            experts = _read_experts_from_sheet(wb[target_name]) + new_experts
            idx = wb.sheetnames.index(target_name)
            del wb[target_name]
            ws = wb.create_sheet(target_name, idx)
            action = "appended"
        else:                                            # different VQ -> new sheet
            # 새 VQ 라운드 = 주제가 다름 → 응답 시트 바로 앞에 '토픽 구분 탭'(빈 시트, 이름만 의미)을 함께 추가.
            # 이름 = spec.project_label(이메일에서 파악한 주제). 기존 '0615(해외 증권사 C, 해외 은행 D) >>' 와 동일 역할.
            label = (spec.get("project_label") or "").strip()
            if label:
                divider_name = _unique_sheet_name(wb, label[:31])
                wb.create_sheet(divider_name)            # append → 다음 응답 시트 직전 위치에 놓임
            target_name = _unique_sheet_name(wb, base_name)
            ws = wb.create_sheet(target_name)
            experts = new_experts
            action = "new_sheet"
    else:                                                # brand-new workbook
        wb = Workbook()
        div = wb.active
        div.title = (spec.get("project_label") or "VQ >>")[:31]
        divider_name = div.title
        target_name = base_name[:31]
        ws = wb.create_sheet(target_name)
        experts = new_experts
        action = "created"

    _write_sheet(ws, questions, experts)
    wb.save(out_path)
    print(json.dumps({"status": "ok", "action": action, "sheet": target_name,
                      "divider": divider_name,
                      "experts_on_sheet": len(experts), "new": len(new_experts),
                      "questions": len(questions), "output": out_path},
                     ensure_ascii=False))


if __name__ == "__main__":
    spec_path, out_path = sys.argv[1], sys.argv[2]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    build(spec, out_path)
